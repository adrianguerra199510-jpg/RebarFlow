"""Espiral suministrada en carrete o con barras individuales (traslapes o acopladores), dividida en sub-marcas."""

from __future__ import annotations

import math
from types import SimpleNamespace

import pytest

from core.element import ElementType, StructuralElement
from core.geometry import GeometryError
from core.params import ParamSpec, param
from core.project import Project
from core.rebar import get_rebar
from core.serialization import project_from_dict, project_to_dict
from core.shape_codes import StraightBar
from core.shape_templates import Spiral
from core.splice import coupler_pieces, lap_pieces, pieces_text, standard_lap_mm
from engine.bbs_calculator import BBSCalculator

N4 = get_rebar("#4")
STOCK = 12_000.0
BIG = {"diameter": 1100.0, "pitch": 100.0, "height": 5600.0}  # ~ 200 m desarrollados


def bbs(shape, *, quantity=1, elements=1, stock=STOCK):
    element = StructuralElement("C-1", ElementType.COLUMN, quantity=elements)
    element.add_group("F1", N4, shape, quantity)
    return BBSCalculator(stock_length_mm=stock).calculate([element])


# ------------------------------------------------------------------ piezas
def test_lap_pieces_add_one_lap_per_joint_and_never_exceed_the_stock_bar():
    for total in (12_001.0, 23_500.0, 24_000.0, 99_999.0, 201_183.0):
        pieces = lap_pieces(total, STOCK, 610.0)
        joints = len(pieces) - 1
        assert sum(pieces) == pytest.approx(total + joints * 610.0)
        assert all(0 < p <= STOCK + 1e-6 for p in pieces)
        assert pieces[-1] > 610.0 - 1e-6  # el resto siempre cubre al menos el traslape
        assert all(p == STOCK for p in pieces[:-1])  # todo lo posible en barras completas


def test_lap_pieces_minimise_bars():
    # 24 m con traslape de 0.6 m: dos barras no alcanzan (23.4 m útiles), hacen falta tres
    assert len(lap_pieces(24_000.0, STOCK, 600.0)) == 3
    assert len(lap_pieces(23_400.0, STOCK, 600.0)) == 2


def test_coupler_pieces_butt_the_bars_without_changing_the_length():
    pieces = coupler_pieces(30_000.0, STOCK)
    assert pieces == (12_000.0, 12_000.0, 6_000.0) and sum(pieces) == 30_000.0
    assert coupler_pieces(24_000.0, STOCK) == (12_000.0, 12_000.0)


def test_short_pieces_are_a_single_bar_without_joints():
    assert lap_pieces(8_000.0, STOCK, 610.0) == (8_000.0,) == coupler_pieces(8_000.0, STOCK)
    assert pieces_text((8_000.0,)) == ""


def test_a_lap_longer_than_half_the_bar_is_rejected():
    with pytest.raises(GeometryError, match="mitad"):
        lap_pieces(30_000.0, STOCK, 7_000.0)


def test_standard_lap_is_48_diameters_but_not_less_than_300_mm():
    assert standard_lap_mm(N4) == pytest.approx(48 * N4.diameter_mm)  # #4: 609.6 mm
    assert standard_lap_mm(get_rebar("#3")) == pytest.approx(48 * get_rebar("#3").diameter_mm)  # 457 mm
    assert standard_lap_mm(SimpleNamespace(diameter_mm=6.0)) == 300.0  # 48 d = 288 mm < 300 mm: manda el mínimo


def test_pieces_text():
    assert pieces_text((12_000.0, 12_000.0, 12_000.0, 4_320.0)) == "3 × 12.00 m + 4.32 m"
    assert pieces_text((12_000.0, 4_320.0)) == "12.00 m + 4.32 m"


# ------------------------------------------------------------------ forma
def test_the_spiral_comes_from_a_coil_by_default_and_does_not_change_its_length():
    shape = Spiral(**BIG)
    supply = shape.supply(N4, STOCK)
    assert supply.coil and supply.parts == () and supply.threads == ()
    assert supply.pieces_mm == (pytest.approx(shape.cut_length(N4).total_mm),)
    assert supply.sub_marks("F1") == ()


def test_a_spiral_from_bars_that_fits_one_bar_has_no_joints_and_no_sub_marks():
    supply = Spiral(300, 75, 300, supply_mode="bars").supply(N4, STOCK)
    assert not supply.coil and len(supply.pieces_mm) == 1 and supply.parts == () and supply.sub_marks("F1") == ()


def test_a_long_spiral_from_bars_is_lapped_by_default_with_the_standard_lap():
    shape = Spiral(**BIG, supply_mode="bars")
    total = shape.cut_length(N4).total_mm
    supply = shape.supply(N4, STOCK)
    joints = len(supply.pieces_mm) - 1
    assert joints == math.ceil((total - 609.6) / (STOCK - 609.6) - 1e-9) - 1
    assert supply.material_mm == pytest.approx(total + joints * standard_lap_mm(N4))
    assert supply.threads == () and supply.warnings == () and f"{joints} traslapes de 610 mm" in supply.notes[0]
    # con traslape las barras no llevan roscas: las completas son una sola sub-marca y el resto, otra
    assert [(p.length_mm, p.count) for p in supply.parts] == [(STOCK, joints), (pytest.approx(supply.pieces_mm[-1]), 1)]
    assert [name for name, _ in supply.sub_marks("F1")] == ["F1.1", "F1.2"]


def test_a_custom_lap_is_used_and_a_short_one_gets_a_warning():
    long_lap = Spiral(**BIG, supply_mode="bars", lap_mm=800.0).supply(N4, STOCK)
    assert long_lap.warnings == () and "traslapes de 800 mm" in long_lap.notes[0]
    short = Spiral(**BIG, supply_mode="bars", lap_mm=200.0).supply(N4, STOCK)
    assert len(short.warnings) == 1 and "200 mm" in short.warnings[0] and "48·d" in short.warnings[0]


def test_a_spiral_with_couplers_keeps_the_length_and_splits_into_first_intermediate_and_last_sub_marks():
    shape = Spiral(**BIG, supply_mode="bars", joint="coupler")  # roscas por omisión: larga + corta
    supply = shape.supply(N4, STOCK)
    bars = len(supply.pieces_mm)
    assert supply.material_mm == pytest.approx(shape.cut_length(N4).total_mm) and supply.warnings == ()
    first, middle, last = supply.parts
    assert (first.count, first.start_thread, first.end_thread) == (1, "", "long")  # la barra que termina lleva la larga
    assert (middle.count, middle.start_thread, middle.end_thread) == (bars - 2, "short", "long")  # las intermedias, idénticas
    assert (last.count, last.start_thread, last.end_thread) == (1, "short", "")  # la que sigue lleva la corta
    assert first.length_mm == middle.length_mm == STOCK and last.length_mm < STOCK
    assert [name for name, _ in supply.sub_marks("F1")] == ["F1.1", "F1.2", "F1.3"]
    assert supply.threads.count("long") == supply.threads.count("short") == bars - 1  # una larga y una corta por unión
    assert f"{bars - 1} acopladores" in supply.notes[0] and "larga en la barra que termina" in supply.notes[0]


def test_two_bars_joined_by_a_coupler_give_exactly_the_two_sub_marks_a1_1_and_a1_2():
    bar = StraightBar(18_000, joint="coupler")
    supply = bar.supply(N4, STOCK)
    assert [(name, p.length_mm, p.start_thread, p.end_thread) for name, p in supply.sub_marks("A1")] == [
        ("A1.1", 12_000.0, "", "long"), ("A1.2", 6_000.0, "short", "")]
    assert supply.threads == ("long", "short")


@pytest.mark.parametrize(("config", "ends"), [("long_short", ("long", "short")), ("short_short", ("short", "short")),
                                              ("long_long", ("long", "long"))])
def test_the_threads_of_the_joint_can_be_long_short_short_short_or_long_long(config, ends):
    supply = StraightBar(18_000, joint="coupler", joint_threads=config).supply(N4, STOCK)
    first, last = supply.parts
    assert (first.end_thread, last.start_thread) == ends and first.start_thread == last.end_thread == ""


def test_a_straight_bar_keeps_its_own_end_threads_on_the_first_and_last_sub_marks():
    supply = StraightBar(30_000, coupler_start="short", coupler_end="long", joint="coupler",
                         joint_threads="short_short").supply(N4, STOCK)
    # primera y segunda barra tienen las mismas roscas (corta-corta): se agrupan en una sola sub-marca
    assert [(p.length_mm, p.count, p.start_thread, p.end_thread) for p in supply.parts] == [
        (12_000.0, 2, "short", "short"), (6_000.0, 1, "short", "long")]


def test_a_straight_bar_that_fits_the_stock_bar_is_not_split():
    supply = StraightBar(12_000, joint="coupler").supply(N4, STOCK)
    assert supply.parts == () and supply.pieces_mm == (12_000.0,) and supply.sub_marks("A1") == ()


def test_the_unused_options_do_not_change_the_result():
    a = Spiral(**BIG, supply_mode="bars", joint="lap", joint_threads="long_long").supply(N4, STOCK)
    b = Spiral(**BIG, supply_mode="bars", joint="lap", joint_threads="short_short").supply(N4, STOCK)
    assert a == b
    assert Spiral(**BIG, supply_mode="coil", joint="coupler").supply(N4, STOCK).coil


def test_unknown_supply_options_are_rejected():
    for bad in ({"supply_mode": "rollo"}, {"joint": "soldadura"}, {"joint_threads": "none"}):
        with pytest.raises(GeometryError):
            Spiral(**BIG, **bad)
    with pytest.raises(GeometryError):
        StraightBar(18_000, joint="soldadura")


def test_param_dependencies():
    specs = {spec.name: spec for spec in Spiral.parameters()}
    assert specs["supply_mode"].depends == () and specs["joint"].active({"supply_mode": "bars"})
    assert not specs["joint"].active({"supply_mode": "coil"})
    assert specs["lap_mm"].active({"supply_mode": "bars", "joint": "lap"})
    assert not specs["lap_mm"].active({"supply_mode": "bars", "joint": "coupler"})
    assert specs["joint_threads"].active({"supply_mode": "bars", "joint": "coupler"})
    assert not specs["joint_threads"].active({"supply_mode": "bars", "joint": None})
    bar = {spec.name: spec for spec in StraightBar.parameters()}
    assert bar["joint"].depends == () and bar["lap_mm"].active({"joint": "lap"}) and not bar["lap_mm"].active({"joint": "coupler"})
    assert bar["joint_threads"].active({"joint": "coupler"})
    assert ParamSpec("x", "x", "length", 1.0).active({})  # sin condiciones, siempre aplica
    assert param("x", depends=(("a", 1),)).metadata["spec"].depends == (("a", 1),)


# ------------------------------------------------------------------ planilla
def test_the_coil_spiral_is_unchanged_in_the_schedule():
    result = bbs(Spiral(**BIG), quantity=2, elements=3)
    line = result.lines[0]
    summary = result.diameter_summaries[0]
    assert line.supply.coil and line.joint_count == 0 and line.pieces_text == "" and line.couplers == ()
    assert line.parts == () and result.thread_rows() == ()
    assert summary.pieces == {} and summary.stock_bars == 0
    assert summary.coil_length_m == pytest.approx(line.total_length_m)


def test_a_lapped_spiral_weighs_and_counts_the_laps_and_enters_the_cutting_plan_by_sub_mark():
    shape = Spiral(**BIG, supply_mode="bars")
    result = bbs(shape, quantity=2, elements=3)  # 6 espirales
    line = result.lines[0]
    summary = result.diameter_summaries[0]
    developed = shape.cut_length(N4).total_mm
    joints = line.joint_count
    assert joints > 10 and not line.supply.coil
    assert line.cut_length_mm == pytest.approx(sum(line.supply.pieces_mm)) and line.cut_length_mm > developed
    assert line.cut_length_mm == pytest.approx(developed + joints * 609.6, abs=40)  # + redondeo de 5 mm por barra
    assert line.total_weight_kg == pytest.approx(6 * N4.mass_for_length(line.cut_length_mm / 1000))
    assert line.total_length_m == pytest.approx(6 * line.cut_length_mm / 1000)
    assert sum(length * count for length, count in summary.pieces.items()) / 1000 == pytest.approx(line.total_length_m)
    assert summary.pieces[STOCK] == 6 * joints  # barras comerciales completas
    assert summary.coil_length_m == 0 and summary.stock_bars >= 6 * joints + 1
    assert summary.marks_for(STOCK) == ("C-1/F1.1",) and "C-1/F1.2" in summary.marks_for(line.parts[1].length_mm)
    assert line.pieces_text.startswith(f"{joints} × 12.00 m + ")
    # los totales salen de las sub-marcas: suman lo mismo que la marca
    assert sum(p.total_weight_kg for p in line.parts) == pytest.approx(line.total_weight_kg)
    assert sum(p.total_length_m for p in line.parts) == pytest.approx(line.total_length_m)
    assert [p.total_bars for p in line.parts] == [6 * joints, 6]
    assert result.total_weight_kg == pytest.approx(line.total_weight_kg)


def test_couplers_and_threads_of_the_joints_are_counted_in_the_summary():
    shape = Spiral(**BIG, supply_mode="bars", joint="coupler")
    result = bbs(shape, quantity=2, elements=3)  # 6 espirales
    line = result.lines[0]
    assert line.cut_length_mm == pytest.approx(shape.cut_length(N4).total_mm, abs=40)
    assert [p.mark for p in line.parts] == ["F1.1", "F1.2", "F1.3"]
    assert [p.count for p in line.parts] == [1, line.joint_count - 1, 1]
    assert line.coupler_text == f"{line.joint_count} × rosca larga + {line.joint_count} × rosca corta"
    counts = {c.kind: c.count for c in result.coupler_counts()}
    assert counts == {"long": 6 * line.joint_count, "short": 6 * line.joint_count}
    (row,) = result.thread_rows()
    assert (row.rebar, row.long, row.short, row.couplers) == (N4, 6 * line.joint_count, 6 * line.joint_count,
                                                              6 * line.joint_count)  # un acoplador junta dos roscas
    assert line.parts[0].thread_text == "rosca larga al final" and line.parts[2].thread_text == "rosca corta al inicio"
    assert line.parts[0].shape.coupler_end == "long" and line.parts[0].shape.coupler_start == "none"


def test_a_straight_bar_with_end_threads_counts_them_together_with_the_joint_threads():
    element = StructuralElement("C-1", ElementType.COLUMN, quantity=2)
    element.add_group("A1", N4, StraightBar(18_000, coupler_start="short", coupler_end="long", joint="coupler"), 3)
    result = BBSCalculator().calculate([element])
    line = result.lines[0]
    assert [(p.mark, p.part.start_thread, p.part.end_thread) for p in line.parts] == [
        ("A1.1", "short", "long"), ("A1.2", "short", "long")]
    counts = {c.kind: c.count for c in result.coupler_counts()}
    assert counts == {"short": 2 * 6, "long": 2 * 6}  # 6 barras: cada una con una larga y una corta (propias + unión)
    (row,) = result.thread_rows()
    assert row.couplers == 12


def test_the_stock_length_of_the_calculator_decides_the_number_of_bars():
    shape = Spiral(**BIG, supply_mode="bars", joint="coupler")
    short = bbs(shape, stock=6_000.0).lines[0]
    long = bbs(shape, stock=12_000.0).lines[0]
    assert short.joint_count > long.joint_count * 1.9


def test_lap_warnings_reach_the_schedule_and_the_detail():
    result = bbs(Spiral(**BIG, supply_mode="bars", lap_mm=200.0))
    assert any("200 mm" in text for _, _, text in result.warnings)
    assert any("Suministro" in note for note in result.lines[0].notes)


def test_a_sub_mark_cannot_collide_with_another_mark_of_the_element():
    element = StructuralElement("C-1", ElementType.COLUMN)
    element.add_group("A1", N4, StraightBar(18_000), 1)
    element.add_group("A1.1", N4, StraightBar(2_000), 1)
    with pytest.raises(ValueError, match="sub-marca A1.1"):
        BBSCalculator().calculate([element])


# ------------------------------------------------------------------ archivos
def test_the_supply_options_survive_save_and_load():
    shape = Spiral(**BIG, supply_mode="bars", joint="coupler", joint_threads="short_short")
    project = Project(name="P")
    element = StructuralElement("C-1", ElementType.COLUMN)
    element.add_group("F1", N4, shape, 1)
    element.add_group("A1", N4, StraightBar(18_000, joint="coupler", joint_threads="long_long", coupler_start="short"), 2)
    project.elements.append(element)
    copy = project_from_dict(project_to_dict(project))
    assert copy.elements[0].groups[0].shape == shape
    assert copy.elements[0].groups[1].shape == element.groups[1].shape
    lapped = Spiral(**BIG, supply_mode="bars", lap_mm=700.0)
    element.groups[0] = type(element.groups[0])("F1", N4, lapped, 1)
    assert project_from_dict(project_to_dict(project)).elements[0].groups[0].shape.lap_mm == 700.0


def test_projects_saved_before_have_a_coil_spiral_and_unsplit_straight_bars():
    project = Project(name="P")
    element = StructuralElement("C-1", ElementType.COLUMN)
    element.add_group("F1", N4, Spiral(**BIG), 1)
    element.add_group("A1", N4, StraightBar(3_000), 1)
    project.elements.append(element)
    data = project_to_dict(project)
    for group in data["elements"][0]["groups"]:
        for key in ("supply_mode", "joint", "lap_mm", "joint_threads"):
            group["shape"]["dimensions"].pop(key, None)
    groups = project_from_dict(data).elements[0].groups
    assert groups[0].shape.supply_mode == "coil" and groups[0].shape.supply(N4, STOCK).coil
    assert groups[1].shape.joint == "lap" and groups[1].shape.supply(N4, STOCK).parts == ()


def test_the_old_joint_coupler_parameter_is_migrated_to_the_threads_of_the_joint():
    from core.shape_codes import create_shape

    for kind, threads in (("short", "short_short"), ("long", "long_long")):
        shape = create_shape("ES", diameter=400, pitch=75, height=2000, supply_mode="bars", joint="coupler",
                             joint_coupler=kind)
        assert shape.joint_threads == threads and not hasattr(shape, "joint_coupler")


# ------------------------------------------------------------------ exportación
def test_exports_list_the_sub_marks_as_rows_with_their_own_weights(tmp_path):
    from openpyxl import load_workbook

    from exporters.excel_exporter import ExcelExporter
    from exporters.pdf_exporter import PdfExporter

    result = bbs(Spiral(**BIG, supply_mode="bars", joint="coupler"), quantity=2)
    line = result.lines[0]
    sheet = load_workbook(ExcelExporter().export(result, tmp_path / "p.xlsx")).worksheets[0]
    rows = {str(row[0].value): row for row in sheet.iter_rows() if row[0].value}
    headers = next([str(c.value) for c in row] for row in sheet.iter_rows() if row[0].value == "Marca")
    weight, length = headers.index("Peso total (kg)"), headers.index("Long. corte (mm)")
    assert {"F1", "F1.1", "F1.2", "F1.3"} <= set(rows)
    assert rows["F1"][weight].value is None and rows["F1"][length].value is None  # la marca deja los pesos a sus sub-marcas
    assert [rows[f"F1.{k}"][length].value for k in (1, 2, 3)] == [p.length_mm for p in line.parts]
    assert sum(rows[f"F1.{k}"][weight].value for k in (1, 2, 3)) == pytest.approx(line.total_weight_kg)
    subtotal = next(row for row in sheet.iter_rows() if str(row[0].value).startswith("Subtotal"))
    assert subtotal[weight].value == pytest.approx(line.total_weight_kg)  # sin contar dos veces
    texts = [str(c.value) for row in sheet.iter_rows() for c in row if c.value is not None]
    assert "Suministro con uniones" not in texts
    pdf = PdfExporter().export(result, tmp_path / "p.pdf")
    assert pdf.stat().st_size > 5_000
