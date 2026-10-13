from pathlib import Path

import pytest

pytest.importorskip("PySide6")

from PySide6.QtCore import QSettings, Qt, QTimer  # noqa: E402
from PySide6.QtWidgets import QApplication, QDialogButtonBox  # noqa: E402

from core.element import ElementType, StructuralElement  # noqa: E402
from core.rebar import get_rebar  # noqa: E402
from core.shape_codes import Stirrup  # noqa: E402
from gui.app import create_application  # noqa: E402
from engine.cutting_report import plan_rows  # noqa: E402
from gui.dialogs import BarDialog, ElementDialog, LegsEditor, ProjectDialog  # noqa: E402
from core.shape_templates import CustomBar, HookedBar  # noqa: E402
from gui import style  # noqa: E402
from gui.main_window import MainWindow  # noqa: E402
from sample_project import build_bridge_project, build_project  # noqa: E402

_REAL_PREFERENCE = (style.preferred_dark, style.remember_dark)  # las funciones reales, antes de aislarlas en cada prueba


@pytest.fixture(scope="session")
def qapp():
    app = create_application([])
    yield app
    # Al terminar la sesión, destruir los widgets antes que la aplicación evita un cierre abrupto del intérprete.
    import gc

    app.closeAllWindows()
    app.processEvents()
    for widget in app.topLevelWidgets():
        widget.deleteLater()
    app.processEvents()
    gc.collect()


@pytest.fixture(autouse=True)
def _isolate_theme(monkeypatch, qapp):
    """Cada prueba empieza en el tema claro y no lee ni escribe la preferencia real del usuario."""
    stored: dict[str, bool] = {}
    monkeypatch.setattr(style, "remember_dark", lambda dark: stored.update(dark=dark))
    monkeypatch.setattr(style, "preferred_dark", lambda: stored.get("dark", False))
    if style.current().dark:  # cambiar el tema repinta todas las ventanas abiertas: solo si hace falta
        style.apply_theme(qapp, False)
    yield stored
    if style.current().dark:
        style.apply_theme(qapp, False)


@pytest.fixture
def window(qapp):
    win = MainWindow(build_project())
    yield win
    for session in win.sessions:  # evita el diálogo de cambios sin guardar al cerrar
        session.modified = False
    win.close()


def ok_enabled(dialog) -> bool:
    return dialog.buttons.button(QDialogButtonBox.StandardButton.Ok).isEnabled()


def select_element(window, name, level):
    root = window.tree.topLevelItem(0)
    for i in range(root.childCount()):
        if root.child(i).text(0) != level:
            continue
        for j in range(root.child(i).childCount()):
            item = root.child(i).child(j)
            if item.text(0).startswith(name):
                window.tree.setCurrentItem(item)
                return window.current_element()
    raise AssertionError(f"{name} / {level} no está en el árbol")


# ------------------------------------------------------------------ ventana principal
def test_window_loads_project_tree_and_tables(window):
    root = window.tree.topLevelItem(0)
    assert root.text(0) == "Edificio de ejemplo"
    assert [root.child(i).text(0) for i in range(root.childCount())] == ["Cimentación", "Nivel 1", "Nivel 2"]
    assert window.current_element().name == "Z-1"
    assert window.bars_table.rowCount() == 3  # 2 marcas + subtotal
    assert window.project_table.rowCount() == 25  # 5 elementos (título + marcas + subtotal) + total general
    assert "Edificio de ejemplo" in window.windowTitle() and not window.windowTitle().startswith("*")
    assert "2,625.2 kg" in window.totals_label.text()
    assert window.act_excel.isEnabled() and window.act_pdf.isEnabled()


def test_every_bar_row_has_a_shape_image(window):
    for row in range(window.project_table.rowCount()):
        if window.project_table.line_for_row(row) is not None:
            icon = window.project_table.item(row, 2).icon()  # columna "Esquema"
            assert not icon.isNull() and not icon.pixmap(160, 60).isNull()


def test_selecting_a_bar_fills_the_detail_panel(window):
    select_element(window, "V-101", "Nivel 1")
    window.bars_table.selectRow(1)
    assert window.current_line().mark == "A2"
    assert "A2" in window.detail.title.text()
    assert "Gancho U" in window.detail.details.text()
    assert not window.detail.preview.pixmap().isNull()
    assert window.act_edit_bar.isEnabled() and window.act_delete_bar.isEnabled()


def test_add_element_and_bars_updates_totals_and_marks_modified(window):
    before = window.result.total_weight_kg
    element = window.add_element_values("C-9", ElementType.COLUMN, "Nivel 3", 2)
    assert window.current_element() is element and window.modified
    assert window.windowTitle().startswith("*")
    assert window.bars_table.rowCount() == 1  # solo el subtotal

    dialog = BarDialog(window, suggested_mark="E1", used_marks=[])
    dialog.shape.setCurrentIndex(dialog.shape.findText("E · Estribo rectangular"))
    dialog._widgets["a"].setValue(320)
    dialog._widgets["b"].setValue(320)
    dialog.quantity.setValue(10)
    dialog.bar.setCurrentIndex(dialog.bar.findData("#3"))
    assert ok_enabled(dialog)
    window.add_group(element, dialog.result_group())
    assert window.result.total_weight_kg == pytest.approx(before + 2 * 10 * 1.43 * 0.560)
    assert window.current_line().mark == "E1"  # la barra nueva queda seleccionada


def test_duplicate_element_is_rejected(window):
    with pytest.raises(ValueError, match="Ya existe"):
        window.add_element_values("V-101", ElementType.BEAM, "Nivel 1", 1)
    window.add_element_values("V-101", ElementType.BEAM, "Nivel 3", 1)  # otro nivel: válido


def test_edit_and_delete_bar_and_element(window):
    element = select_element(window, "V-101", "Nivel 1")
    old = element.groups[0]
    dialog = BarDialog(window, group=old, used_marks=[g.mark for g in element.groups if g is not old])
    dialog.quantity.setValue(5)
    window.replace_group(element, old, dialog.result_group())
    assert element.groups[0].quantity == 5 and window.current_line().mark == "A1"

    window.delete_group(element, element.groups[0])
    assert [g.mark for g in element.groups] == ["A2", "A3", "E1"]

    n = len(window.project.elements)
    window.delete_element_object(element)
    assert len(window.project.elements) == n - 1
    assert window.current_element() is not None  # se selecciona otro elemento


def test_update_element_values(window):
    element = select_element(window, "V-101", "Nivel 1")
    window.update_element_values(element, "V-102", ElementType.BEAM, "Nivel 1", 6)
    assert element.name == "V-102" and element.quantity == 6
    with pytest.raises(ValueError):
        window.update_element_values(element, "C-1", ElementType.BEAM, "Nivel 1", 1)


def test_calculation_error_is_shown_not_raised(window):
    window.project.stock_length_mm = 3000  # las barras de 5.6 m ya no caben
    window.refresh()
    assert window.result is None and "excede la barra comercial" in window.error
    assert not window.banner.isHidden() and "No se puede calcular" in window.banner.text()
    assert not window.act_excel.isEnabled() and not window.act_pdf.isEnabled()
    with pytest.raises(ValueError):
        window.write_pdf(Path("x.pdf"))
    window.project.stock_length_mm = 12000
    window.refresh()
    assert window.result is not None and window.banner.isHidden()


def test_save_open_and_export(window, tmp_path, qapp):
    window.add_element_values("C-9", ElementType.COLUMN, "Nivel 3", 2)
    assert window.modified
    path = tmp_path / "obra.rebarflow"
    window.write_project(path)
    assert not window.modified and window.path == path

    other = MainWindow()
    assert other.open_path(path)
    assert [e.name for e in other.project.elements] == [e.name for e in window.project.elements]
    assert other.result.total_weight_kg == pytest.approx(window.result.total_weight_kg)
    other.close()

    assert window.write_excel(tmp_path / "planilla.xlsx").stat().st_size > 5000
    assert window.write_pdf(tmp_path / "planilla.pdf").read_bytes().startswith(b"%PDF")


def test_open_invalid_file_reports_failure(window, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QMessageBox
    shown = []
    monkeypatch.setattr(QMessageBox, "critical", lambda *a, **k: shown.append(a[2]))
    bad = tmp_path / "mal.rebarflow"
    bad.write_text("{}", encoding="utf-8")
    assert window.open_path(bad) is False
    assert shown and "mal.rebarflow" in shown[0]


def test_empty_project_window(qapp):
    win = MainWindow()
    assert win.current_element() is None
    assert not win.act_excel.isEnabled() and not win.act_add_bar.isEnabled()
    assert win.totals_label.text() == "Sin barras"
    win.close()


# ------------------------------------------------------------------ diálogos
def test_bar_dialog_rejects_impossible_geometry_and_recovers(qapp):
    dialog = BarDialog(None, suggested_mark="B1")
    dialog.shape.setCurrentIndex(dialog.shape.findText("L · Bastón L"))
    dialog.bar.setCurrentIndex(dialog.bar.findData("#8"))
    dialog._widgets["a"].setValue(60)
    assert not ok_enabled(dialog) and "mínimo geométrico" in dialog.info.text()
    with pytest.raises(ValueError):
        dialog.result_group()
    dialog._widgets["a"].setValue(1000)
    assert ok_enabled(dialog) and "Longitud de corte" in dialog.info.text()
    assert dialog.result_group().shape.a == 1000


def test_bar_dialog_mark_and_stock_validation(qapp):
    dialog = BarDialog(None, suggested_mark="A1", used_marks=["A1"])
    assert not ok_enabled(dialog) and "ya existe" in dialog.info.text()
    dialog.mark.setText("A2")
    assert ok_enabled(dialog)
    dialog.mark.setText("  ")
    assert not ok_enabled(dialog)
    dialog.mark.setText("A2")
    dialog.shape.setCurrentIndex(dialog.shape.findText("L · Bastón L"))  # una barra doblada no se divide sola
    dialog._widgets["a"].setValue(13000)
    assert not ok_enabled(dialog) and "excede la barra comercial" in dialog.info.text()
    dialog.shape.setCurrentIndex(dialog.shape.findText("R · Recta"))  # una recta larga sí: se divide en sub-marcas
    dialog._widgets["length"].setValue(13000)
    assert ok_enabled(dialog) and "A2.1" in dialog.info.text() and "A2.2" in dialog.info.text()


def test_bar_dialog_builds_fields_from_shape_and_keeps_values(qapp):
    dialog = BarDialog(None)
    dialog.shape.setCurrentIndex(dialog.shape.findText("U · Gancho U"))
    assert list(dialog._widgets) == ["a", "b", "c"]
    dialog._widgets["a"].setValue(1234)
    dialog.shape.setCurrentIndex(dialog.shape.findText("E · Estribo rectangular"))
    assert list(dialog._widgets) == ["a", "b", "hook_start", "ext_start", "hook_end", "ext_end"]
    assert dialog._widgets["a"].value() == 1234  # se conserva el valor del campo de igual nombre
    assert dialog._widgets["ext_start"].value() == 0 and dialog._widgets["ext_end"].value() == 0  # "estándar"
    assert dialog.result_group().shape.ext_start is None
    dialog._widgets["ext_start"].setValue(50)
    assert "75" in dialog.info.text() and "Aviso" in dialog.info.text()  # gancho por debajo del mínimo
    assert dialog.result_group().shape == Stirrup(1234, 400, ext_start=50)


def _choose(dialog, name, value):
    combo = dialog._widgets[name]
    combo.setCurrentIndex(combo.findData(value))


def test_bar_dialog_offers_bar_14(qapp):
    dialog = BarDialog(None)
    designations = [dialog.bar.itemData(i) for i in range(dialog.bar.count())]
    assert designations[-1] == "#14" and designations[:-1] == [f"#{n}" for n in range(3, 12)]
    dialog.bar.setCurrentIndex(dialog.bar.findData("#14"))
    dialog.shape.setCurrentIndex(dialog.shape.findText("L · Bastón L"))
    assert ok_enabled(dialog) and "Longitud de corte" in dialog.info.text()
    dialog.shape.setCurrentIndex(dialog.shape.findText("E · Estribo rectangular"))
    assert not ok_enabled(dialog) and "mayores a #8" in dialog.info.text()  # la norma no define estribos de #14


def test_bar_dialog_stirrup_has_a_hook_per_leg(qapp):
    dialog = BarDialog(None)
    dialog.shape.setCurrentIndex(dialog.shape.findText("E · Estribo rectangular"))
    assert not any(dialog.shape.itemText(i).startswith("E90") for i in range(dialog.shape.count()))
    _choose(dialog, "hook_end", 90)
    shape = dialog.result_group().shape
    assert (shape.hook_start, shape.hook_end) == (135, 90)
    assert "135° · 75  /  90° · 75" in dialog.info.text()


def test_bar_dialog_ring_with_hooks_and_no_lap(qapp):
    dialog = BarDialog(None)
    dialog.shape.setCurrentIndex(dialog.shape.findText("AR · Aro circular"))
    assert list(dialog._widgets) == ["diameter", "lap", "hook_start", "ext_start", "hook_end", "ext_end"]
    dialog._widgets["lap"].setValue(0)  # los ganchos cierran el aro: sin traslape
    _choose(dialog, "hook_start", 135)
    _choose(dialog, "hook_end", 135)
    shape = dialog.result_group().shape
    assert shape.lap == 0 and shape.hook_start == shape.hook_end == 135
    assert "2 × 135° · 75" in dialog.info.text()


def test_bar_dialog_spiral_extra_turns_and_hooks(qapp):
    dialog = BarDialog(None)
    dialog.shape.setCurrentIndex(dialog.shape.findText("ES · Espiral"))
    assert list(dialog._widgets) == ["diameter", "pitch", "height", "turns_start", "turns_end", "hook_start", "ext_start",
                                     "hook_end", "ext_end", "supply_mode", "joint", "lap_mm", "joint_threads"]
    dialog._widgets["turns_start"].setValue(2.0)
    dialog._widgets["turns_end"].setValue(0.5)
    _choose(dialog, "hook_end", 90)
    shape = dialog.result_group().shape
    assert (shape.turns_start, shape.turns_end, shape.hook_start, shape.hook_end) == (2.0, 0.5, 0, 90)
    assert "2 al inicio y 0.5 al final" in dialog.info.text() and "90° · 75" in dialog.info.text()


def test_bar_dialog_couplers_on_free_ends_only(qapp):
    dialog = BarDialog(None)
    dialog.shape.setCurrentIndex(dialog.shape.findText("R · Recta"))
    assert list(dialog._widgets) == ["length", "coupler_start", "coupler_end", "joint", "lap_mm", "joint_threads"]
    assert [dialog._widgets["coupler_end"].itemText(i) for i in range(3)] == ["Sin rosca", "Rosca corta", "Rosca larga"]
    _choose(dialog, "coupler_start", "short")
    _choose(dialog, "coupler_end", "long")
    shape = dialog.result_group().shape
    assert (shape.coupler_start, shape.coupler_end) == ("short", "long")
    assert "Roscas para acoplador: rosca corta + rosca larga" in dialog.info.text() and ok_enabled(dialog)
    dialog.shape.setCurrentIndex(dialog.shape.findText("G · Barra con ganchos"))  # los ganchos por defecto (90°) chocan
    assert not ok_enabled(dialog) and "termina en gancho" in dialog.info.text()
    _choose(dialog, "hook_start", 0)
    _choose(dialog, "hook_end", 0)
    assert ok_enabled(dialog)
    dialog.shape.setCurrentIndex(dialog.shape.findText("L · Bastón L"))
    assert [name for name in dialog._widgets if name.startswith("coupler")] == ["coupler_start"]  # la pata es el gancho


def test_bar_dialog_prefills_existing_group(qapp):
    group = build_project().elements[2].groups[3]  # estribo E1 de V-101
    dialog = BarDialog(None, group=group, used_marks=["A1", "A2", "A3"])
    assert dialog.mark.text() == "E1" and dialog.bar.currentData() == "#3"
    assert dialog.shape.currentData() is Stirrup
    assert dialog._widgets["a"].value() == 220 and dialog._widgets["b"].value() == 420
    assert dialog.quantity.value() == 38 and ok_enabled(dialog)
    assert dialog.result_group().shape == group.shape


def test_element_dialog(qapp):
    dialog = ElementDialog(None, levels=["Nivel 1", "Nivel 2"])
    assert not ok_enabled(dialog)  # falta el nombre
    dialog.name.setText("V-1")
    assert ok_enabled(dialog)
    dialog.type.setCurrentIndex(dialog.type.findData(ElementType.FOOTING))
    dialog.quantity.setValue(3)
    assert dialog.values() == ("V-1", ElementType.FOOTING, "Nivel 2", 3, "")  # propone el último nivel


def test_project_dialog_applies_values(qapp):
    project = build_project()
    dialog = ProjectDialog(None, project)
    dialog.name.setText("Torre Norte")
    dialog.stock.setValue(9000)
    dialog.rounding.setValue(10)
    dialog.paper.setCurrentIndex(dialog.paper.findData("a4"))
    dialog.apply_to(project)
    assert (project.name, project.stock_length_mm, project.length_rounding_mm, project.paper) == ("Torre Norte", 9000, 10, "a4")


# ------------------------------------------------------------------ confirmaciones en español
def click_button(label: str, texts: list[str]) -> None:
    """Programa el clic en el botón ``label`` del diálogo modal que esté abierto."""
    def click():
        box = QApplication.activeModalWidget()
        texts.extend(b.text() for b in box.buttons())
        next(b for b in box.buttons() if b.text() == label).click()
    QTimer.singleShot(100, click)


def test_yes_no_dialog_uses_spanish_buttons(window):
    texts: list[str] = []
    click_button("Sí", texts)
    assert window._ask("Título", "¿Seguro?") is True
    assert texts == ["Sí", "No"]
    click_button("No", [])
    assert window._ask("Título", "¿Seguro?") is False


def test_discard_changes_dialog(window):
    assert window._confirm_discard() is True  # sin cambios no pregunta
    window.modified = True
    texts: list[str] = []
    click_button("Descartar", texts)
    assert window._confirm_discard() is True
    assert sorted(texts) == ["Cancelar", "Descartar", "Guardar"]  # Qt los ordena por rol
    click_button("Cancelar", [])
    assert window._confirm_discard() is False
    window.modified = False


def test_delete_slots_ask_for_confirmation(window, monkeypatch):
    element = select_element(window, "V-101", "Nivel 1")
    window.bars_table.selectRow(0)
    monkeypatch.setattr(MainWindow, "_ask", lambda self, title, text: False)
    window.delete_bar()
    assert len(element.groups) == 4  # respondió "No": no se borra
    monkeypatch.setattr(MainWindow, "_ask", lambda self, title, text: True)
    window.delete_bar()
    assert len(element.groups) == 3
    window.delete_element()
    assert all(e.name != "V-101" or e.level != "Nivel 1" for e in window.project.elements)


# ------------------------------------------------------------------ plan de corte, ganchos, plantillas
def test_plan_tab_lists_the_cutting_sequence(window):
    assert window.tabs.tabText(3) == "Plan de corte"
    table = window.plan_table
    picture_rows = [r for r in range(table.rowCount()) if table.item(r, 3) is not None and not table.item(r, 3).icon().isNull()]
    assert len(picture_rows) == sum(len(plan_rows(s)) for s in window.result.diameter_summaries)
    assert any("plan lineal" in (table.item(r, 0).text() if table.item(r, 0) else "") for r in range(table.rowCount()))
    window.project.stock_length_mm = 3000
    window.refresh()
    assert table.rowCount() == 1 and "Aún no hay barras" in table.item(0, 0).text()  # sin resultado, mensaje vacío


def test_tables_have_no_diameter_shape_hook_or_remark_columns(window):
    for table in (window.bars_table, window.project_table):
        headers = [table.horizontalHeaderItem(c).text().replace("\n", " ") for c in range(table.columnCount())]
        assert headers[:3] == ["Marca", "Barra", "Esquema"]
        for removed in ("Ø (mm)", "Forma", "Ganchos (° · mm)", "Observaciones", "A", "B", "C", "D", "E"):
            assert removed not in headers
        assert len(headers) == 9  # Marca, Barra, Esquema + los seis totales
        assert headers[-1] == "Peso total (kg)"
    # los ganchos siguen en el dibujo (rótulo G) y en el panel de detalle de la barra
    select_element(window, "V-101", "Nivel 1")
    window.bars_table.selectRow(1)  # A2: U con ganchos de 90°
    assert "Ganchos (° · mm)" in window.detail.details.text() and "90°" in window.detail.details.text()


def test_row_with_a_code_warning_is_marked_red_with_a_tooltip(window):
    select_element(window, "V-101", "Nivel 1")
    element = window.current_element()
    from core.rebar import get_rebar
    element.add_group("E9", get_rebar("#3"), Stirrup(220, 420, ext_start=50, ext_end=50), 2)
    window.refresh()
    table = window.bars_table
    row = table.row_of("E9")
    assert table.item(row, 0).foreground().color().name().lower() == "#c00000"
    assert "mínimo normativo" in table.item(row, 0).toolTip()
    assert table.item(table.row_of("E1"), 0).toolTip() == ""


def test_bar_dialog_offers_every_template_with_a_thumbnail(qapp):
    dialog = BarDialog(None, suggested_mark="T")
    classes = [dialog.shape.itemData(i) for i in range(dialog.shape.count())]
    assert {"R", "L", "U", "E", "G", "EU", "Z", "Q", "O", "ET", "ETR", "AR", "ES", "X"} <= {c.code for c in classes}
    for index, cls in enumerate(classes):
        assert not dialog.shape.itemIcon(index).isNull(), cls.code
        dialog.shape.setCurrentIndex(index)
        assert dialog.shape_hint.text() == cls.description
        assert ok_enabled(dialog), f"la plantilla {cls.code} debe ser válida con sus valores por defecto"
        assert dialog.result_group().shape == cls.template()


def test_hooked_bar_dialog_validates_hooks_and_style(qapp):
    dialog = BarDialog(None, suggested_mark="G1")
    dialog.shape.setCurrentIndex(dialog.shape.findText("G · Barra con ganchos"))
    start = dialog._widgets["hook_start"]
    start.setCurrentIndex(start.findData(135))
    assert not ok_enabled(dialog) and "no hay gancho estándar de 135°" in dialog.info.text()
    style = dialog._widgets["style"]
    style.setCurrentIndex(style.findData("stirrup"))
    assert ok_enabled(dialog) and "Ganchos (° · mm)" in dialog.info.text()
    assert dialog.result_group().shape.hook_start == 135 and dialog.result_group().cut.hooks[0].extension_mm == 75.0
    dialog._widgets["ext_start"].setValue(50)  # bajo el mínimo de la norma: aviso, no error
    assert ok_enabled(dialog) and "Aviso" in dialog.info.text()


def test_legs_editor_operations(qapp):
    editor = LegsEditor((300.0, 200.0, 300.0), (90.0, -90.0))
    assert editor.values() == ((300.0, 200.0, 300.0), (90.0, -90.0))
    editor.add_leg()
    assert editor.values() == ((300.0, 200.0, 300.0, 300.0), (90.0, -90.0, 90.0))
    editor.reverse()
    assert editor.values() == ((300.0, 300.0, 200.0, 300.0), (-90.0, 90.0, -90.0))
    editor.mirror()
    assert editor.values()[1] == (90.0, -90.0, 90.0)
    editor.table.selectRow(1)
    editor.remove_leg()  # desaparece el tramo 2 y el giro con que salía de él
    assert editor.values() == ((300.0, 200.0, 300.0), (90.0, 90.0))
    editor.table.setCurrentCell(-1, -1)
    editor.remove_leg()  # sin selección quita el último tramo
    assert editor.values() == ((300.0, 200.0), (90.0,))
    editor.remove_leg()
    editor.remove_leg()
    assert len(editor.values()[0]) == 1 and editor.values()[1] == ()  # siempre queda al menos un tramo


def test_custom_shape_through_the_dialog(qapp):
    dialog = BarDialog(None, suggested_mark="X1")
    dialog.shape.setCurrentIndex(dialog.shape.findText("X · Personalizada"))
    assert dialog._editor is not None
    dialog._editor.set_values((400.0, 300.0, 200.0), (90.0, -45.0))
    hook = dialog._widgets["hook_end"]
    hook.setCurrentIndex(hook.findData(90))
    assert ok_enabled(dialog)
    shape = dialog.result_group().shape
    assert shape == CustomBar(legs=(400.0, 300.0, 200.0), angles=(90.0, -45.0), hook_end=90)
    dialog._editor.set_values((400.0, 300.0), (200.0,))  # un giro de 200° es imposible
    assert not ok_enabled(dialog)


def test_dialog_prefills_existing_hooked_and_custom_shapes(qapp):
    for original in (HookedBar(1200, 135, 100, 90, None, 1, "stirrup"),
                     CustomBar(legs=(300.0, 200.0, 300.0), angles=(90.0, -90.0), hook_start=90)):
        group = build_project().elements[2].groups[0].__class__("P1", get_rebar("#4"), original, 2)
        dialog = BarDialog(None, group=group)
        assert dialog.shape.currentData() is type(original)
        assert ok_enabled(dialog) and dialog.result_group().shape == original


def test_other_element_type_and_project_remnant_setting(qapp, window):
    dialog = ElementDialog(None, levels=["Nivel 1"])
    dialog.name.setText("MURETE-1")
    dialog.type.setCurrentIndex(dialog.type.findData(ElementType.OTHER))
    assert dialog.type.currentText() == "Otros" and dialog.values()[1] is ElementType.OTHER
    element = window.add_element_values(*dialog.values())
    assert element.element_type is ElementType.OTHER and window.current_element() is element
    assert "Otros" in window.tree.currentItem().text(0)  # el árbol muestra «MURETE-1 · Otros ×1»
    assert window.result is not None  # el cálculo y las tablas funcionan con el tipo nuevo

    project = build_project()
    pdialog = ProjectDialog(None, project)
    assert pdialog.reusable.value() == 1000
    pdialog.reusable.setValue(2500)
    pdialog.apply_to(project)
    assert project.reusable_remnant_mm == 2500


def test_open_catalog_example_from_the_menu(window):
    window.open_catalog()
    assert window.project.name == "Catálogo de formas" and not window.modified
    assert window.bars_table.rowCount() == 36  # 26 marcas + 9 sub-marcas (R4, ES3, ES4) + subtotal
    pictures = [r for r in range(window.bars_table.rowCount()) if window.bars_table.line_for_row(r)]
    assert all(not window.bars_table.item(r, 2).icon().isNull() for r in pictures)
    assert window.plan_table.rowCount() > 1 and window.act_catalog.text().startswith("Abrir ejemplo")


def test_closed_detail_panel_can_be_reopened(window):
    dock = window.detail_dock
    assert not dock.isHidden()
    dock.close()  # la X del panel
    assert dock.isHidden()
    window.act_detail.trigger()  # menú Ver → Detalle de barra (también el botón de la barra y F4)
    assert not dock.isHidden() and window.act_detail.isChecked()
    window.act_detail.trigger()
    assert dock.isHidden() and not window.act_detail.isChecked()
    menus = {action.text().replace("&", ""): action.menu() for action in window.menuBar().actions()}
    assert window.act_detail in menus["Ver"].actions()
    assert window.act_detail.shortcut().toString() == "F4"


# ------------------------------------------------------------------ modo nocturno
def test_theme_button_sits_in_the_top_right_corner_and_switches_the_whole_interface(window, qapp, _isolate_theme):
    button = window.theme_button
    assert window.menuBar().cornerWidget(Qt.Corner.TopRightCorner) is button
    assert button.defaultAction() is window.act_theme and window.act_theme.text() == "Modo nocturno"
    assert not style.current().dark and "#F4F6FA" in qapp.styleSheet()
    menus = {action.text().replace("&", ""): action.menu() for action in window.menuBar().actions()}
    assert window.act_theme in menus["Ver"].actions()

    button.click()  # claro → nocturno
    assert style.current() is style.DARK and window.act_theme.text() == "Modo diurno"
    assert style.DARK.window in qapp.styleSheet() and qapp.palette().window().color().name().lower() == style.DARK.window.lower()
    assert _isolate_theme == {"dark": True}  # se recuerda para la próxima vez
    # tablas, panel de detalle y esquemas se redibujan con los colores del tema
    assert window.project_table.item(0, 0).background().color().name().lower() == style.DARK.band.lower()
    window.bars_table.selectRow(1)
    assert style.DARK.muted in window.detail.details.text()
    dark_icon = window.bars_table.item(1, 2).icon().pixmap(176, 66).toImage()

    button.click()  # nocturno → claro
    assert style.current() is style.LIGHT and window.act_theme.text() == "Modo nocturno"
    assert _isolate_theme == {"dark": False}
    assert window.project_table.item(0, 0).background().color().name().lower() == style.LIGHT.band.lower()
    window.bars_table.selectRow(1)
    light_icon = window.bars_table.item(1, 2).icon().pixmap(176, 66).toImage()
    assert dark_icon != light_icon


def test_theme_keeps_the_selected_bar_and_can_skip_remembering(window, _isolate_theme):
    select_element(window, "V-101", "Nivel 1")
    window.bars_table.selectRow(2)
    mark = window.current_line().mark
    window.set_dark_mode(True, remember=False)
    assert window.current_line().mark == mark and _isolate_theme == {}
    window.toggle_theme()
    assert not style.current().dark and window.current_line().mark == mark


def test_dialogs_use_the_theme_colors(window):
    window.set_dark_mode(True, remember=False)
    dialog = BarDialog(None, suggested_mark="E7")
    dialog.shape.setCurrentIndex(dialog.shape.findText("E · Estribo rectangular"))
    dialog._widgets["ext_start"].setValue(50)  # aviso de norma
    assert style.DARK.muted in dialog.info.text() and style.DARK.hook in dialog.info.text() and style.DARK.red in dialog.info.text()
    assert "#595959" not in dialog.info.text()
    assert not dialog.preview.pixmap().isNull()


def test_theme_preference_is_stored_with_qsettings(tmp_path, monkeypatch):
    preferred, remember = _REAL_PREFERENCE
    monkeypatch.setattr(style, "_settings", lambda: QSettings(str(tmp_path / "rebarflow.ini"), QSettings.Format.IniFormat))
    assert preferred() is False  # claro por omisión
    remember(True)
    assert preferred() is True and "nocturno" in (tmp_path / "rebarflow.ini").read_text(encoding="utf-8")
    remember(False)
    assert preferred() is False


def test_application_starts_in_the_remembered_theme(qapp, _isolate_theme):
    _isolate_theme["dark"] = True
    create_application([])
    assert style.current().dark
    create_application([], dark=False)  # un tema explícito manda sobre la preferencia
    assert not style.current().dark


def test_theme_icons_are_drawn_for_both_directions(qapp):
    from gui.icon import theme_icon

    moon, sun = theme_icon(to_dark=True), theme_icon(to_dark=False)
    assert not moon.isNull() and not sun.isNull()
    assert moon.pixmap(32, 32).toImage() != sun.pixmap(32, 32).toImage()


def test_light_is_the_default_even_if_windows_uses_its_dark_theme(qapp, tmp_path, monkeypatch, _isolate_theme):
    preferred, _ = _REAL_PREFERENCE
    monkeypatch.setattr(style, "_settings", lambda: QSettings(str(tmp_path / "vacio.ini"), QSettings.Format.IniFormat))
    assert preferred() is False  # sin preferencia guardada: claro
    qapp.styleHints().setColorScheme(Qt.ColorScheme.Dark)  # un equipo con el tema oscuro de Windows
    create_application([])  # arranque normal: lee la preferencia (ninguna → claro)
    palette = qapp.palette()
    assert not style.current().dark
    assert palette.window().color().name().lower() == style.LIGHT.window.lower()
    assert palette.base().color().name().lower() == "#ffffff" and palette.text().color().name().lower() == style.LIGHT.text.lower()
    scheme = qapp.styleHints().colorScheme()
    assert scheme in (Qt.ColorScheme.Light, Qt.ColorScheme.Unknown)  # «Unknown» solo en la plataforma de pruebas sin pantalla
    win = MainWindow(build_project())
    assert win.act_theme.text() == "Modo nocturno"  # el botón ofrece pasar a nocturno, no al revés
    win.modified = False
    win.close()


# ------------------------------------------------------------------ obra sin niveles (puente), encabezado y logo
def tree_paths(window):
    paths = []

    def walk(item, prefix):
        for i in range(item.childCount()):
            child = item.child(i)
            paths.append(prefix + (child.text(0),))
            walk(child, prefix + (child.text(0),))

    walk(window.tree.topLevelItem(0), ())
    return paths


def test_bridge_example_groups_by_structure_and_substructure_without_levels(window):
    window.open_bridge()
    assert window.project.name == "Paso inferior de ejemplo" and not window.modified
    paths = tree_paths(window)
    assert ("E-1",) in paths and ("E-1", "GENERICA") in paths
    assert ("E-1", "GENERICA", "Losa de transición 1  ·  Losa  ×2") in paths
    assert ("P-1", "Pila", "Viga cabezal  ·  Viga  ×1") in paths and ("P-1", "Cimentación", "Zapata de pila  ·  Zapata  ×1") in paths
    root = window.tree.topLevelItem(0)
    assert [root.child(i).text(0) for i in range(root.childCount())] == ["E-1", "P-1", "Barrera de concreto  ·  Muro  ×4"]
    assert root.child(2).childCount() == 0  # el elemento sin estructura cuelga directamente del proyecto
    assert window.current_element().name == "Losa de transición 1"
    assert "E-1 / GENERICA" in window.element_title.text()
    menus = {action.text().replace("&", ""): action.menu() for action in window.menuBar().actions()}
    assert window.act_bridge in menus["Archivo"].actions()


def test_element_without_levels_can_be_added_and_edited(window):
    window.open_bridge()
    element = window.add_element_values("Muro de contención", ElementType.WALL)
    assert element.level == element.substructure == "" and element.location == ""
    root = window.tree.topLevelItem(0)
    assert root.child(root.childCount() - 1).text(0).startswith("Muro de contención")
    window.update_element_values(element, "Muro de contención", ElementType.WALL, "E-2", 3, "Estribo")
    assert element.path == ("E-2", "Estribo") and element.quantity == 3
    with pytest.raises(ValueError, match="Ya existe el elemento 'Muro de contención' en 'E-2 / Estribo'"):
        window.add_element_values("Muro de contención", ElementType.WALL, "E-2", 1, "Estribo")
    window.add_element_values("Muro de contención", ElementType.WALL, "E-2", 1, "Otro")  # otra subestructura: válido
    with pytest.raises(ValueError, match="Ya existe el elemento 'Barrera de concreto'."):
        window.add_element_values("Barrera de concreto", ElementType.WALL)


def test_element_dialog_levels_are_optional_and_the_first_group_takes_the_project_name(qapp):
    dialog = ElementDialog(None, levels=["E-1", "P-1"], substructures=["GENERICA"], group_label="Estructura")
    assert dialog.level.lineEdit().placeholderText() == "(sin estructura)"
    assert dialog.substructure.lineEdit().placeholderText() == "(sin subestructura)"
    assert [dialog.substructure.itemText(i) for i in range(dialog.substructure.count())] == ["GENERICA"]
    assert not ok_enabled(dialog)  # falta el nombre
    dialog.name.setText("Barrera")
    dialog.level.setCurrentText("")
    assert ok_enabled(dialog)  # sin nivel ni subestructura es válido
    assert dialog.values() == ("Barrera", ElementType("Viga"), "", 1, "")
    dialog.level.setCurrentText("P-1")
    dialog.substructure.setCurrentText("Pila")
    assert dialog.values()[2:] == ("P-1", 1, "Pila")
    existing = ElementDialog(None, StructuralElement("X", ElementType.SLAB, "E-1", 2, substructure="GENERICA"))
    assert existing.level.currentText() == "E-1" and existing.substructure.currentText() == "GENERICA"


def test_project_dialog_edits_header_data_group_label_and_logo(qapp, tmp_path):
    from exporters.logo import LogoError

    project = build_bridge_project()
    dialog = ProjectDialog(None, project)
    assert dialog.client.text() == "Consorcio de ejemplo" and dialog.group_label.currentText() == "Estructura"
    assert dialog.logo() == b"" and not dialog.logo_clear.isEnabled()
    dialog.client.setText("Otro cliente")
    dialog.work_code.setText("12345")
    dialog.group_label.setCurrentText("Tramo")
    png = tmp_path / "logo.png"
    from PIL import Image

    Image.new("RGB", (800, 300), (10, 90, 100)).save(png)
    dialog.load_logo(png)
    assert dialog.logo().startswith(b"\x89PNG") and dialog.logo_clear.isEnabled() and not dialog.logo_preview.pixmap().isNull()
    dialog.apply_to(project)
    assert (project.client, project.work_code, project.group_label) == ("Otro cliente", "12345", "Tramo")
    assert project.logo_png == dialog.logo()
    bad = tmp_path / "no_es_imagen.png"
    bad.write_bytes(b"texto")
    with pytest.raises(LogoError):
        dialog.load_logo(bad)
    assert project.logo_png == dialog.logo()  # un archivo no válido no cambia el logo
    dialog.logo_clear.click()
    assert dialog.logo() == b"" and dialog.logo_preview.text() == "sin logo"


def test_window_exports_carry_the_header_data_and_logo(window, tmp_path):
    from openpyxl import load_workbook

    from exporters.logo import normalize_logo

    window.open_bridge()
    window.project.logo_png = normalize_logo(_png_bytes())
    sheet = load_workbook(window.write_excel(tmp_path / "p.xlsx"))["Planilla BBS"]
    texts = [str(sheet.cell(r, 2).value) for r in range(1, 12)]
    assert any("Cliente: Consorcio de ejemplo" in t and "Cód. obra: 00001" in t for t in texts)
    assert any(image.anchor._from.row == 0 for image in sheet._images)  # el logo, arriba a la izquierda
    window.project.group_label = "Tramo"
    window.refresh()
    assert window.result.group_label == "Tramo" and window.result.level_title("") == "Sin tramo"


def _png_bytes() -> bytes:
    from io import BytesIO

    from PIL import Image

    buffer = BytesIO()
    Image.new("RGB", (300, 100), (20, 100, 110)).save(buffer, format="PNG")
    return buffer.getvalue()


def test_bar_dialog_spiral_supply_options_follow_the_choices(qapp):
    dialog = BarDialog(None, stock_length_mm=12_000.0)
    dialog.shape.setCurrentIndex(dialog.shape.findText("ES · Espiral"))
    enabled = lambda: [name for name in ("supply_mode", "joint", "lap_mm", "joint_threads") if dialog._widgets[name].isEnabled()]  # noqa: E731
    assert enabled() == ["supply_mode"]  # carrete: no hay uniones
    assert "Suministro: carrete" in dialog.info.text() and "barras:" not in dialog.info.text()
    _choose(dialog, "supply_mode", "bars")
    assert enabled() == ["supply_mode", "joint", "lap_mm"]  # traslape por omisión
    _choose(dialog, "joint", "coupler")
    assert enabled() == ["supply_mode", "joint", "joint_threads"]
    assert [dialog._widgets["joint_threads"].itemText(i) for i in range(3)] == ["Larga + corta", "Corta + corta", "Larga + larga"]
    _choose(dialog, "joint_threads", "long_long")
    shape = dialog.result_group().shape
    assert (shape.supply_mode, shape.joint, shape.joint_threads) == ("bars", "coupler", "long_long")
    text = dialog.info.text()
    assert "barras:" in text and "A1.1" in text and "A1.2" in text and "rosca larga al final" in text
    assert "Roscas para acoplador: " in text and "rosca larga" in text
    _choose(dialog, "joint", "lap")
    dialog._widgets["lap_mm"].setValue(500.0)
    assert dialog.result_group().shape.lap_mm == 500.0 and "traslapes de 500 mm" in dialog.info.text()
    assert "Roscas para acoplador: " not in dialog.info.text() and ok_enabled(dialog)
    _choose(dialog, "supply_mode", "coil")
    assert "Suministro: carrete" in dialog.info.text()


def test_bar_dialog_straight_bar_joint_options_and_sub_marks(qapp):
    dialog = BarDialog(None, suggested_mark="A1", stock_length_mm=12_000.0)
    dialog.shape.setCurrentIndex(dialog.shape.findText("R · Recta"))
    names = ("joint", "lap_mm", "joint_threads")
    assert list(dialog._widgets) == ["length", "coupler_start", "coupler_end", *names]
    enabled = lambda: [n for n in names if dialog._widgets[n].isEnabled()]  # noqa: E731
    assert enabled() == ["joint", "lap_mm"]
    dialog._widgets["length"].setValue(30_000)
    _choose(dialog, "joint", "coupler")
    assert enabled() == ["joint", "joint_threads"]
    text = dialog.info.text()
    assert "A1.1" in text and "A1.2" in text and "A1.3" in text and "rosca corta al inicio, rosca larga al final" in text
    assert ok_enabled(dialog)
    group = dialog.result_group()
    assert group.shape.joint == "coupler" and group.shape.length == 30_000


def test_bar_dialog_spiral_from_bars_changes_the_steel_with_laps(qapp):
    dialog = BarDialog(None, stock_length_mm=12_000.0)
    dialog.shape.setCurrentIndex(dialog.shape.findText("ES · Espiral"))
    coil_weight = dialog.info.text().split("peso unitario <b>")[1].split(" kg")[0]
    _choose(dialog, "supply_mode", "bars")
    bars_weight = dialog.info.text().split("peso unitario <b>")[1].split(" kg")[0]
    assert float(bars_weight.replace(",", "")) > float(coil_weight.replace(",", ""))  # los traslapes consumen acero


def test_theme_button_keeps_room_for_its_longest_text_even_with_a_big_font(window, qapp):
    """En Windows el texto salía recortado («Modo …turno»): el botón reserva el ancho del texto más largo."""
    from PySide6.QtGui import QFont

    button = window.theme_button
    original = qapp.font()
    try:
        for point_size in (original.pointSize(), 14, 18):
            font = QFont(original)
            font.setPointSize(point_size)
            qapp.setFont(font)
            widths = []
            for dark in (False, True):
                window.set_dark_mode(dark, remember=False)
                metrics = button.fontMetrics()
                longest = max(metrics.horizontalAdvance(t) for t in ("Modo nocturno", "Modo diurno"))
                hint = button.sizeHint()
                assert hint.width() >= longest + button.iconSize().width() + 20, (point_size, dark)
                assert button.minimumSizeHint() == hint
                widths.append(hint.width())
            assert abs(widths[0] - widths[1]) <= 2  # el ancho no cambia al cambiar el texto
    finally:
        qapp.setFont(original)
        window.set_dark_mode(False, remember=False)


# ------------------------------------------------------------------ sub-marcas en las tablas y el detalle
def _window_with_split_bars(window):
    from core.shape_codes import StraightBar
    from core.shape_templates import Spiral

    element = StructuralElement("C-1", ElementType.COLUMN, "Nivel 3", 2)
    element.add_group("A1", get_rebar("#5"), StraightBar(30_000, joint="coupler"), 3)
    element.add_group("F1", get_rebar("#4"), Spiral(1100, 100, 5600, supply_mode="bars", joint="coupler"), 1)
    element.add_group("B1", get_rebar("#4"), StraightBar(2_000), 2)
    window.project.elements.append(element)
    window.refresh(select=element)
    return element


def test_the_bars_table_lists_the_sub_marks_under_their_mark_with_their_own_numbers(window):
    element = _window_with_split_bars(window)
    table = window.bars_table
    marks = [table.item(r, 0).text() if table.item(r, 0) else "" for r in range(table.rowCount())]
    assert marks[:8] == ["A1", "A1.1", "A1.2", "A1.3", "F1", "F1.1", "F1.2", "F1.3"] and marks[8] == "B1"
    headers = [table.horizontalHeaderItem(c).text().replace("\n", " ") for c in range(table.columnCount())]
    peso = headers.index("Peso total (kg)")
    assert table.item(0, peso).text() == "" and table.item(1, peso).text() != ""  # la marca deja los pesos a sus sub-marcas
    schedule = next(s for s in window.result.element_schedules if s.element is element)
    shown = sum(float(table.item(r, peso).text().replace(",", "")) for r in range(table.rowCount())
                if table.item(r, 0) and table.item(r, peso) and table.item(r, peso).text() and "Subtotal" not in table.item(r, 0).text())
    assert shown == pytest.approx(schedule.total_weight_kg, abs=0.05)
    assert table.part_for_row(0) is None and table.part_for_row(1).mark == "A1.1"
    assert table.line_for_row(2) is table.line_for_row(0)  # la sub-fila pertenece a la marca A1
    assert table.row_of("A1") == 0 and table.row_of("F1") == 4


def test_selecting_a_sub_mark_shows_that_bar_in_the_detail_and_edit_acts_on_the_whole_mark(window, monkeypatch):
    _window_with_split_bars(window)
    window.bars_table.selectRow(2)  # A1.2: la barra intermedia
    assert window.current_part().mark == "A1.2" and window.current_line().mark == "A1"
    title, details = window.detail.title.text(), window.detail.details.text()
    assert "A1.2" in title and "Sub-marca de" in details and "rosca corta al inicio, rosca larga al final" in details
    assert "A1.1 ×1" in details and "se arma con" in details
    window.bars_table.selectRow(0)
    assert window.current_part() is None and "Sub-marcas" in window.detail.details.text()
    opened = []

    def fake_exec(self):
        opened.append(self.mark.text())
        return False

    monkeypatch.setattr(BarDialog, "exec", fake_exec)
    window.bars_table.selectRow(3)  # A1.3
    window.edit_bar()
    assert opened == ["A1"] and window.act_edit_bar.isEnabled()  # se edita la marca completa


def test_the_project_table_and_the_summary_show_sub_marks_acopladores_and_roscas(window):
    _window_with_split_bars(window)
    window.tabs.setCurrentIndex(1)
    table = window.project_table
    rows = [table.item(r, 0).text() for r in range(table.rowCount()) if table.item(r, 0)]
    assert {"A1.1", "A1.2", "A1.3", "F1.1", "F1.2", "F1.3"} <= set(rows)
    row = next(r for r in range(table.rowCount()) if table.item(r, 0) and table.item(r, 0).text() == "F1.2")
    table.selectRow(row)
    assert "F1.2" in window.detail.title.text()
    from PySide6.QtWidgets import QLabel

    titles = [label.text() for label in window.summary.findChildren(QLabel)]
    assert any("Acopladores y roscas por barra" in t for t in titles)


def test_the_sub_rows_of_a_spiral_draw_spirals_and_the_detail_panel_shows_that_spiral(window):
    from core.shape_templates import Spiral

    _window_with_split_bars(window)
    table = window.bars_table
    spiral_row = table.row_of("F1")
    assert type(table.line_for_row(spiral_row).shape) is Spiral
    parts = [table.part_for_row(spiral_row + k) for k in (1, 2, 3)]
    assert [type(p.shape) for p in parts] == [Spiral, Spiral, Spiral]
    table.selectRow(spiral_row + 2)  # F1.2: la espiral intermedia
    details = window.detail.details.text()
    assert "F1.2" in window.detail.title.text() and "rosca corta al inicio, rosca larga al final" in details
    assert window.detail.preview.pixmap() is not None and not window.detail.preview.pixmap().isNull()
