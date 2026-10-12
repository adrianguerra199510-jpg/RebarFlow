import math

import pytest

from core.geometry import GeometryError
from core.rebar import BendUse, HookType, get_rebar
from core.serialization import project_from_dict, project_to_dict
from core.project import Project
from core.element import ElementType, StructuralElement
from core.shape_codes import (
    LBar,
    Stirrup,
    StraightBar,
    UBar,
    available_shapes,
    create_shape,
)
from core.shape_templates import (
    CrankBar,
    CustomBar,
    HookedBar,
    OmegaBar,
    RingTie,
    Spiral,
    TrapStirrup,
    TriStirrup,
    UStirrup,
    ZBar,
)
from engine.bbs_calculator import BBSCalculator

N3, N4, N5, N6, N9, N14 = (get_rebar(n) for n in ("#3", "#4", "#5", "#6", "#9", "#14"))


def delta90(rebar, use=BendUse.MAIN):
    r, d = rebar.bend_radius_mm(use), rebar.diameter_mm
    return 2 * (r + d) - math.pi / 2 * (r + d / 2)


def walk(items, heading=0.0):
    """Recorre una barra con rectas y arcos y devuelve (longitud, ancho, alto) del eje."""
    x = y = 0.0
    heading = math.radians(heading)
    length, xs, ys = 0.0, [0.0], [0.0]
    for item in items:
        if item[0] == "s":
            x, y = x + item[1] * math.cos(heading), y + item[1] * math.sin(heading)
            length += item[1]
            xs.append(x), ys.append(y)
        else:  # ("t", grados con signo, radio)
            turn, radius = math.radians(item[1]), item[2]
            sign = 1 if turn > 0 else -1
            steps = 200
            for k in range(steps):
                h0 = heading + turn * k / steps
                h1 = heading + turn * (k + 1) / steps
                x += sign * radius * (math.sin(h1) - math.sin(h0))
                y += -sign * radius * (math.cos(h1) - math.cos(h0))
                xs.append(x), ys.append(y)
            length += abs(turn) * radius
            heading += turn
    return length, max(xs) - min(xs), max(ys) - min(ys)


# ------------------------------------------------------------------ catálogo
def test_catalog_contains_every_requested_typology():
    codes = set(available_shapes())
    assert {"R", "L", "U", "E", "G", "EU", "Z", "Q", "O", "ET", "ETR", "AR", "ES", "X"} <= codes


@pytest.mark.parametrize("code", ["R", "L", "U", "E", "G", "EU", "Z", "Q", "O", "ET", "ETR", "AR", "ES", "X"])
def test_every_template_computes_draws_and_describes_itself(code):
    cls = available_shapes()[code]
    assert cls.name and cls.description and cls.parameters()
    shape = cls.template()
    for rebar in (N3, N4, N5):
        cut = shape.cut_length(rebar)
        assert cut.total_mm > 0 and math.isfinite(cut.total_mm)
        sketch = shape.sketch(rebar)
        assert sketch.paths and all(len(path) >= 2 for path in sketch.paths)
    assert shape.sketch(None).paths  # sin barra también dibuja
    assert shape.dimensions()


def test_templates_either_compute_or_raise_geometry_error_for_large_bars():
    for cls in available_shapes().values():
        for designation in ("#6", "#8", "#10"):
            try:
                cls.template().cut_length(get_rebar(designation))
            except GeometryError:
                pass  # una plantilla de 300 mm puede ser imposible con #10: debe avisar, no fallar


# ------------------------------------------------------------------ motor de ganchos
def test_hooked_bar_with_two_90_hooks_equals_u_bar():
    ext = N4.standard_hook(HookType.DEG_90).extension_mm
    hooked = HookedBar(a=1000, hook_start=90, hook_end=90).cut_length(N4)
    u = UBar.with_standard_hooks(N4, a=1000).cut_length(N4)
    assert hooked.total_mm == pytest.approx(u.total_mm)
    assert [h.extension_mm for h in hooked.hooks] == pytest.approx([ext, ext])


@pytest.mark.parametrize(("start", "end"), [(90, 135), (135, 135), (180, 90), (180, 180), (135, 0), (0, 90)])
def test_hooked_bar_length_matches_independent_walk(start, end):
    shape = HookedBar(a=1200, hook_start=start, hook_end=end, style="stirrup" if 135 in (start, end) else "main")
    use = BendUse.STIRRUP if shape.style == "stirrup" else BendUse.MAIN
    rebar = N4
    d, r, R = rebar.diameter_mm, rebar.bend_radius_mm(use), rebar.centerline_radius_mm(use)
    cut = shape.cut_length(rebar)
    ext = {h.angle_deg: h.extension_mm for h in cut.hooks}
    items = []
    extent_start = extent_end = 0.0
    if start:
        items += [("s", ext[start]), ("t", start, R)]
        extent_start = R + d / 2 if start >= 90 else 0
    # tramo recto: cota exterior menos lo que consume cada gancho
    straight = 1200 - (r + d if start else 0) - (r + d if end else 0)
    items.append(("s", straight))
    if end:
        items += [("t", end, R), ("s", ext[end])]
    length, _, _ = walk(items)
    assert cut.total_mm == pytest.approx(length, abs=1e-6)


def test_u_stirrup_geometry_closes_on_its_outer_dimensions():
    """Con ganchos de 135° el eje encierra exactamente a − d de ancho y b − d de alto (+ d/2 de arco)."""
    rebar = N3
    shape = UStirrup(a=300, b=400, hook_start=135, hook_end=135)
    d, r, R = rebar.diameter_mm, rebar.bend_radius_mm(BendUse.STIRRUP), rebar.centerline_radius_mm(BendUse.STIRRUP)
    cut = shape.cut_length(rebar)
    ext = cut.hooks[0].extension_mm
    leg = 400 - (r + d) - (r + d)  # cota exterior menos el extremo del gancho y la tangente del doblez inferior
    base = 300 - 2 * (r + d)
    # Recorrido: cola → gancho (giro a la izquierda) → pata → doblez → base → doblez → pata → gancho → cola
    items = [("s", ext), ("t", 135, R), ("s", leg), ("t", 90, R), ("s", base), ("t", 90, R), ("s", leg), ("t", 135, R), ("s", ext)]
    length, _, _ = walk(items, heading=45.0)
    assert cut.total_mm == pytest.approx(length, abs=1e-6)
    # Caja de la parte en U (sin las colas): ancho exterior = a, alto exterior = b
    u_only = [("s", leg), ("t", 90, R), ("s", base), ("t", 90, R), ("s", leg)]
    _, width, height = walk(u_only, heading=-90.0)
    assert width + d == pytest.approx(300 - 0, abs=1e-6)
    assert height + (r + d) + d / 2 == pytest.approx(400, abs=1e-6)  # + extremo exterior del gancho (r + d)


def test_standard_hook_lengths_are_reported():
    cut = LBar.with_standard_hook(N4, a=1000).cut_length(N4)
    assert len(cut.hooks) == 1 and cut.hooks[0].angle_deg == 90
    assert cut.hooks[0].extension_mm == pytest.approx(152.4)  # 12 d
    assert cut.hooks[0].standard_mm == pytest.approx(152.4)
    assert cut.hook_text == "90° · 152"
    stirrup = Stirrup(220, 420).cut_length(N4)
    assert stirrup.hook_text == "2 × 135° · 76"
    assert Stirrup(220, 420, hook_start=90, hook_end=90).cut_length(N4).hook_text == "2 × 90° · 76"


def test_hook_extension_validation():
    with pytest.raises(GeometryError, match="no hay gancho estándar de 135°"):
        HookedBar(a=1000, hook_start=135, hook_end=0).cut_length(N4)  # estilo barra: ACI no define 135°
    forced = HookedBar(a=1000, hook_start=135, hook_end=0, ext_start=120).cut_length(N4)
    assert forced.hooks[0].extension_mm == 120 and "sin extensión estándar" in forced.warnings[0]
    short = HookedBar(a=1000, hook_start=90, hook_end=0, ext_start=50).cut_length(N4)
    assert "mínimo normativo" in short.warnings[0]
    assert HookedBar(a=1000, hook_start=0, hook_end=0).cut_length(N4).total_mm == 1000


def test_rectangular_stirrup_90_formula():
    d, R = N3.diameter_mm, N3.centerline_radius_mm(BendUse.STIRRUP)
    ext = max(6 * d, 75)
    cut = Stirrup(220, 420, hook_start=90, hook_end=90).cut_length(N3)
    expected = 2 * (220 - d) + 2 * (420 - d) - 8 * R + (3 * math.pi / 2 + math.pi) * R + 2 * ext
    assert cut.total_mm == pytest.approx(expected)


# ------------------------------------------------------------------ polilíneas
def test_z_bar_and_omega_are_sum_of_dimensions_minus_bend_deductions():
    assert ZBar(250, 200, 250).cut_length(N4).total_mm == pytest.approx(700 - 2 * delta90(N4))
    omega = OmegaBar(200, 300, 400, 300, 200).cut_length(N4)
    assert omega.total_mm == pytest.approx(1400 - 4 * delta90(N4))
    assert omega.hooks == ()


def test_crank_bar_uses_the_bend_deduction_of_its_angle():
    r, d = N5.bend_radius_mm(), N5.diameter_mm
    angle = 45.0
    deduction = 2 * (r + d) * math.tan(math.radians(angle) / 2) - math.radians(angle) * (r + d / 2)
    cut = CrankBar(600, 300, 600, angle).cut_length(N5)
    assert cut.total_mm == pytest.approx(1500 - 2 * deduction)
    assert any("Desnivel exterior" in note for note in cut.notes)
    with pytest.raises(GeometryError):
        CrankBar(600, 300, 600, angle=95)  # ángulo fuera de rango


def test_geometry_errors_name_the_leg():
    with pytest.raises(GeometryError, match="cota B"):
        OmegaBar(200, 60, 400, 300, 200).cut_length(N6)


# ------------------------------------------------------------------ polígonos
def test_trapezoid_with_equal_bases_is_the_rectangular_stirrup():
    trap = TrapStirrup(bottom=220, top=220, height=420, hook_start=135, hook_end=135).cut_length(N3)
    rect = Stirrup(220, 420).cut_length(N3)
    assert trap.total_mm == pytest.approx(rect.total_mm)


def test_trapezoid_and_triangle_geometry():
    trap = TrapStirrup(bottom=400, top=250, height=300)
    assert trap.leg() == pytest.approx(math.hypot(300, 75))
    assert trap.base_angle() == pytest.approx(math.degrees(math.atan2(300, 75)))
    tri = TriStirrup(300, 300, 300)
    assert tri.interior_angles() == pytest.approx((60.0, 60.0, 60.0))
    assert sum(180 - a for a in tri.interior_angles()) == pytest.approx(360.0)
    with pytest.raises(GeometryError, match="no forman un triángulo"):
        TriStirrup(100, 100, 250)


def test_equilateral_triangle_stirrup_length():
    rebar, s = N3, 300.0
    r, d, R = rebar.bend_radius_mm(BendUse.STIRRUP), rebar.diameter_mm, rebar.centerline_radius_mm(BendUse.STIRRUP)
    ext = max(6 * d, 75)
    turn = math.radians(120)
    straight = s - (r + d) * 2 * math.tan(turn / 2)
    expected = 3 * straight + 2 * turn * R + 2 * math.radians(135) * R + 2 * ext
    assert TriStirrup(s, s, s, hook_start=135, hook_end=135).cut_length(rebar).total_mm == pytest.approx(expected)
    notes = TriStirrup(s, s, s).cut_length(rebar).notes[0]
    assert "60°" in notes


# ------------------------------------------------------------------ aro y espiral
def test_ring_matches_bs8666_code_75():
    """BS 8666 código 75: L = π·(A − d) + B."""
    cut = RingTie(diameter=400, lap=250).cut_length(N3)
    assert cut.total_mm == pytest.approx(math.pi * (400 - N3.diameter_mm) + 250)
    with pytest.raises(GeometryError, match="mínimo geométrico"):
        RingTie(diameter=40, lap=100).cut_length(N3)


def test_spiral_length_and_spacing_warning():
    d = N4.diameter_mm
    shape = Spiral(diameter=400, pitch=75, height=2000, turns_start=1.5, turns_end=1.5)
    n = (2000 - d) / 75
    centre = math.pi * (400 - d)
    assert shape.cut_length(N4).total_mm == pytest.approx(n * math.hypot(centre, 75) + 2 * 1.5 * centre)
    assert not shape.cut_length(N4).warnings
    assert "25–75" in Spiral(400, 30, 2000).cut_length(N4).warnings[0]
    assert "25–75" in Spiral(400, 120, 2000).cut_length(N4).warnings[0]
    with pytest.raises(GeometryError):
        Spiral(400, 10, 2000).cut_length(N4)  # paso menor que el diámetro de la barra


# ------------------------------------------------------------------ BS 8666
def test_l_bar_agrees_with_bs8666_code_11():
    """BS 8666 código 11: L = A + B − 0.5r − d. Con r = 3d coincide con la deducción exacta (2.50d)."""
    r, d = N4.bend_radius_mm(), N4.diameter_mm
    bs = 1000 + 300 - 0.5 * r - d
    assert LBar(1000, 300).cut_length(N4).total_mm == pytest.approx(bs, abs=0.01 * d)


# ------------------------------------------------------------------ forma personalizada
def test_custom_bar_matches_equivalent_template():
    custom = CustomBar(legs=(250, 200, 250), angles=(90, -90)).cut_length(N4)
    assert custom.total_mm == pytest.approx(ZBar(250, 200, 250).cut_length(N4).total_mm)
    with_hooks = CustomBar(legs=(1000,), angles=(), hook_start=90, hook_end=90).cut_length(N4)
    assert with_hooks.total_mm == pytest.approx(HookedBar(1000, 90, None, 90, None).cut_length(N4).total_mm)


def test_custom_bar_accepts_lists_and_validates():
    shape = CustomBar(legs=[300, 200, 300], angles=[60, -60])  # listas (p. ej. venidas de JSON) → tuplas
    assert shape.legs == (300.0, 200.0, 300.0) and shape.angles == (60.0, -60.0)
    hash(shape)  # sigue siendo inmutable y hashable
    with pytest.raises(GeometryError, match="ángulos de giro"):
        CustomBar(legs=(300, 200), angles=(90, 90))
    with pytest.raises(GeometryError):
        CustomBar(legs=(300, 200), angles=(0,))
    with pytest.raises(GeometryError):
        CustomBar(legs=(300, 200), angles=(180,))
    with pytest.raises(GeometryError):
        CustomBar(legs=(300, -5), angles=(90,))
    assert list(CustomBar(legs=(1, 2, 3, 4), angles=(90, 90, 90)).dimensions()) == ["A", "B", "C", "D"]


def test_create_shape_and_json_round_trip_of_new_shapes():
    assert create_shape("Q", a=500, b=250, c=500, angle=30).angle == 30
    project = Project(elements=[StructuralElement("X-1", ElementType.OTHER, "Nivel 1", 2)])
    element = project.elements[0]
    element.add_group("1", N4, CustomBar(legs=(400, 300, 200), angles=(90, -45), hook_end=90), 3)
    element.add_group("2", N3, TriStirrup(300, 280, 260, hook_start=90, hook_end=90), 10)
    element.add_group("3", N4, Spiral(350, 70, 1800), 2)
    element.add_group("4", N4, HookedBar(1000, 135, 100, 90, None, 1, "stirrup"), 4)
    loaded = project_from_dict(project_to_dict(project))
    for original, copy in zip(element.groups, loaded.elements[0].groups):
        assert copy.shape == original.shape
        assert copy.cut_length_mm == pytest.approx(original.cut_length_mm)
    assert loaded.elements[0].element_type is ElementType.OTHER
    # y el cálculo del proyecto completo funciona con todas
    assert BBSCalculator().calculate(loaded.elements).total_weight_kg > 0


def test_parameters_are_described_for_forms():
    for cls in available_shapes().values():
        names = [spec.name for spec in cls.parameters()]
        assert len(names) == len(set(names))
        for spec in cls.parameters():
            assert spec.label
            if spec.kind == "choice":
                assert spec.choices
    custom = {spec.name: spec for spec in CustomBar.parameters()}
    assert custom["legs"].kind == "lengths" and custom["angles"].hidden


# ------------------------------------------------------------------ barra #14
def test_bar_14_cut_length_with_standard_hook():
    leg = N14.standard_hook(HookType.DEG_90).outside_leg_mm  # 516 + 215 + 43
    assert leg == pytest.approx(774.0)
    cut = LBar.with_standard_hook(N14, a=1500).cut_length(N14)
    assert cut.total_mm == pytest.approx(1500 + 774 - delta90(N14))
    assert cut.hook_text == "90° · 516"
    with pytest.raises(GeometryError, match="mayores a #8"):
        Stirrup(600, 600).cut_length(N14)


# ------------------------------------------------------------------ estribo único con patas independientes
@pytest.mark.parametrize("start,end", [(135, 135), (90, 90), (135, 90), (90, 135)])
def test_stirrup_legs_formula_for_every_hook_combination(start, end):
    d, R = N3.diameter_mm, N3.centerline_radius_mm(BendUse.STIRRUP)
    ext = max(6 * d, 75)
    cut = Stirrup(220, 420, hook_start=start, hook_end=end).cut_length(N3)
    expected = 2 * (220 - d) + 2 * (420 - d) - 8 * R + (3 * math.pi / 2 + math.radians(start + end)) * R + 2 * ext
    assert cut.total_mm == pytest.approx(expected)
    assert [h.angle_deg for h in cut.hooks] == [start, end]


def test_stirrup_legs_have_independent_extensions():
    cut = Stirrup(220, 420, hook_start=135, hook_end=90, ext_start=80, ext_end=130).cut_length(N3)
    assert [h.extension_mm for h in cut.hooks] == [80, 130]
    assert cut.hook_text == "135° · 80  /  90° · 130"
    swapped = Stirrup(220, 420, hook_start=90, hook_end=135, ext_start=130, ext_end=80).cut_length(N3)
    assert swapped.total_mm == pytest.approx(cut.total_mm)  # el largo no depende de cuál pata es la inicial
    assert len(Stirrup(220, 420, ext_start=50, ext_end=50).cut_length(N3).warnings) == 1  # un solo aviso, no dos


@pytest.mark.parametrize("shape", [Stirrup(), UStirrup(), TriStirrup(), TrapStirrup()])
def test_every_stirrup_family_takes_a_hook_on_each_leg(shape):
    assert [p.name for p in type(shape).parameters()][-4:] == ["hook_start", "ext_start", "hook_end", "ext_end"]
    mixed = type(shape)(**{**{p.name: getattr(shape, p.name) for p in type(shape).parameters()},
                           "hook_start": 135, "hook_end": 90})
    assert [h.angle_deg for h in mixed.cut_length(N3).hooks] == [135, 90]


def test_there_is_a_single_stirrup_template():
    assert "E90" not in available_shapes() and available_shapes()["E"] is Stirrup


def test_legacy_stirrup_and_spiral_parameters_still_load():
    assert create_shape("E90", a=220, b=420, hook_extension_mm=None) == Stirrup(220, 420, hook_start=90, hook_end=90)
    assert create_shape("E", a=220, b=420, hook_extension_mm=90) == Stirrup(220, 420, ext_start=90, ext_end=90)
    assert create_shape("EU", a=300, b=400, hook=90, hook_extension_mm=None) == UStirrup(300, 400, hook_start=90, hook_end=90)
    assert create_shape("ET", a=300, b=300, c=300, hook=135, hook_extension_mm=None) == TriStirrup(300, 300, 300)
    assert create_shape("ES", diameter=400, pitch=75, height=2000, extra_turns=2.0) == Spiral(
        400, 75, 2000, turns_start=2.0, turns_end=2.0)


# ------------------------------------------------------------------ aro y espiral con ganchos
def test_ring_hooks_add_their_arc_and_extension():
    d, R = N3.diameter_mm, N3.centerline_radius_mm(BendUse.STIRRUP)
    base = math.pi * (400 - d)
    assert RingTie(400, lap=0).cut_length(N3).total_mm == pytest.approx(base)
    both = RingTie(400, lap=0, hook_start=135, hook_end=135).cut_length(N3)
    assert both.total_mm == pytest.approx(base + 2 * (math.radians(135) * R + 75))
    assert both.hook_text == "2 × 135° · 75"
    mixed = RingTie(400, lap=100, hook_start=135, hook_end=90, ext_end=120).cut_length(N3)
    assert mixed.total_mm == pytest.approx(base + 100 + (math.radians(135) + math.radians(90)) * R + 75 + 120)
    assert RingTie(400, lap=250, hook_end=90).cut_length(N3).hook_text == "90° · 75"
    assert RingTie(400, lap=250).cut_length(N3).hooks == ()  # sin ganchos sigue siendo el código 75 de BS 8666


def test_ring_hook_of_big_bars_needs_an_explicit_extension():
    with pytest.raises(GeometryError, match="no hay gancho estándar de 135°"):
        RingTie(900, lap=0, hook_start=135).cut_length(N9)
    forced = RingTie(900, lap=0, hook_start=135, ext_start=200).cut_length(N9)
    assert forced.hooks[0].extension_mm == 200 and "sin extensión estándar" in forced.warnings[0]


def test_spiral_has_independent_extra_turns_and_hooks_at_each_end():
    d, R = N4.diameter_mm, N4.centerline_radius_mm(BendUse.STIRRUP)
    n, centre = (2000 - d) / 75, math.pi * (400 - d)
    working = n * math.hypot(centre, 75)
    uneven = Spiral(400, 75, 2000, turns_start=2.0, turns_end=0.5).cut_length(N4)
    assert uneven.total_mm == pytest.approx(working + 2.5 * centre)
    assert Spiral(400, 75, 2000, turns_start=0.5, turns_end=2.0).cut_length(N4).total_mm == pytest.approx(uneven.total_mm)
    assert Spiral(400, 75, 2000, turns_start=0, turns_end=0).cut_length(N4).total_mm == pytest.approx(working)
    hooked = Spiral(400, 75, 2000, turns_start=0, turns_end=0, hook_start=135, hook_end=90, ext_end=150).cut_length(N4)
    assert hooked.total_mm == pytest.approx(working + (math.radians(135) + math.radians(90)) * R + max(6 * d, 75) + 150)
    assert hooked.hook_text == "135° · 76  /  90° · 150"
    with pytest.raises(GeometryError):
        Spiral(400, 75, 2000, turns_start=-1)


# ------------------------------------------------------------------ referencias de cota en el dibujo
def dim_texts(shape, rebar=None):
    return {label.text for label in shape.sketch(rebar).labels if label.kind == "dim"}


def hook_texts(shape, rebar):
    return [label.text for label in shape.sketch(rebar).labels if label.kind == "hook"]


@pytest.mark.parametrize("code", sorted(available_shapes()))
def test_every_dimension_label_carries_the_table_reference(code):
    shape = available_shapes()[code].template()
    dims = shape.dimensions()
    texts = dim_texts(shape, N3)
    assert texts
    for text in texts:
        name, value = text.split("=")
        assert name in dims and value == f"{dims[name]:.0f}", text
    assert {text.split("=")[0] for text in texts} == set(dims)  # todas las columnas de la planilla están en el dibujo


def test_labels_follow_the_dimension_letters_not_the_drawing_order():
    assert dim_texts(UBar(600, 300, 350)) == {"A=600", "B=300", "C=350"}  # la base es A aunque se dibuje en medio
    assert dim_texts(Stirrup(220, 420)) == {"A=220", "B=420"}  # los lados iguales se rotulan una vez
    assert dim_texts(UStirrup(300, 400)) == {"A=300", "B=400"}
    trap = TrapStirrup(400, 250, 300)
    assert dim_texts(trap) == {"A=400", "B=250", "C=300", "D=309"}
    assert len(trap.sketch().guides) == 3  # la altura se acota con una línea de cota punteada
    assert dim_texts(CustomBar(legs=(1, 2, 3, 4), angles=(90, 90, 90))) == {"A=1", "B=2", "C=3", "D=4"}


def test_hook_labels_show_the_angle_and_the_extension():
    # ganchos iguales: un solo rótulo; distintos: uno por gancho (el del final junto al último lado, el del inicio abajo)
    assert hook_texts(Stirrup(220, 420), N3) == ["G 135° · 75"]
    assert hook_texts(Stirrup(220, 420, hook_start=90, hook_end=90), N3) == ["G 90° · 75"]
    assert hook_texts(Stirrup(220, 420, hook_start=135, hook_end=90), N3) == ["G 90° · 75", "G 135° · 75"]
    assert hook_texts(Stirrup(220, 420, ext_end=120), N3) == ["G 135° · 120", "G 135° · 75"]
    assert hook_texts(TriStirrup(300, 300, 300, hook_start=90), N3) == ["G 135° · 75", "G 90° · 75"]
    assert hook_texts(RingTie(400, lap=0, hook_start=135, hook_end=135), N3) == ["G 135° · 75"]
    assert hook_texts(RingTie(400, lap=250, hook_end=90), N3) == ["G 90° · 75"]
    assert hook_texts(RingTie(400, lap=250, hook_start=135, hook_end=90), N3) == ["G 135° · 75", "G 90° · 75"]
    assert hook_texts(RingTie(400, lap=250), N3) == []
    assert hook_texts(Spiral(400, 75, 2000, hook_start=135, hook_end=90, ext_end=150), N4) == ["G 135° · 76", "G 90° · 150"]
    assert hook_texts(Spiral(400, 75, 2000), N4) == []
    # barras: gancho explícito en cada punta y la pata de una L o de una U (doblez de 90°)
    assert hook_texts(HookedBar(a=1800, hook_start=90, hook_end=180), N4) == ["G 90° · 152", "G 180° · 65"]
    assert hook_texts(HookedBar(a=1800, hook_start=135, hook_end=0, ext_start=120), N4) == ["G 135° · 120"]
    tail = f"{300 - (N5.bend_radius_mm() + N5.diameter_mm):.0f}"  # la pata recta: cota exterior menos la tangente del doblez
    assert hook_texts(LBar(1200, 300), N5) == [f"G 90° · {tail}"]
    assert hook_texts(UBar(600, 300, 300), N5) == [f"G 90° · {tail}", f"G 90° · {tail}"]
    assert hook_texts(Stirrup(220, 420), None) == []  # sin barra no se conoce el largo del gancho


@pytest.mark.parametrize("code", sorted(available_shapes()))
@pytest.mark.parametrize("size", [(320, 120), (640, 240)])
def test_sketch_image_keeps_every_label_inside_the_picture(code, size):
    """Ningún texto, cota punteada ni rótulo toca el borde de la imagen (solo la barra puede llegar hasta él)."""
    from exporters.shape_images import BAR_COLOR, render_sketch

    shape = available_shapes()[code].template()
    image = render_sketch(shape.sketch(N3), *size)
    width, height = image.size
    border = [image.getpixel((x, y)) for x in range(width) for y in (0, height - 1)]
    border += [image.getpixel((x, y)) for y in range(height) for x in (0, width - 1)]
    not_bar = [p for p in border if p[3] > 60 and sum((a - b) ** 2 for a, b in zip(p[:3], BAR_COLOR[:3])) > 25**2]
    assert not not_bar


# ------------------------------------------------------------------ couplers
def test_coupler_on_a_free_end_does_not_change_the_cut_length():
    plain = StraightBar(3000).cut_length(N5)
    both = StraightBar(3000, coupler_start="short", coupler_end="long").cut_length(N5)
    assert both.total_mm == plain.total_mm == 3000
    assert plain.couplers == () and plain.coupler_text == ""
    assert both.couplers == ("short", "long") and both.coupler_text == "rosca corta + rosca larga"
    assert StraightBar(3000, coupler_end="short").cut_length(N5).coupler_text == "rosca corta"
    assert StraightBar(3000, coupler_start="long", coupler_end="long").cut_length(N5).coupler_text == "2 × rosca larga"


@pytest.mark.parametrize("shape", [
    ZBar(coupler_start="short", coupler_end="long"), CrankBar(coupler_end="long"),
    OmegaBar(coupler_start="long", coupler_end="short"),
    CustomBar(legs=(400, 300), angles=(90,), coupler_start="short"), LBar(1200, 300, coupler_start="long"),
])
def test_every_open_bar_template_takes_couplers(shape):
    cut = shape.cut_length(N5)
    assert cut.couplers and shape.sketch(N5).boxes
    without = type(shape)(**{p.name: getattr(shape, p.name) for p in type(shape).parameters()
                             if not p.name.startswith("coupler")})
    assert cut.total_mm == pytest.approx(without.cut_length(N5).total_mm)  # el coupler es un accesorio


def test_coupler_is_not_allowed_on_an_end_that_finishes_in_a_hook():
    with pytest.raises(GeometryError, match="extremo inicial termina en gancho"):
        HookedBar(a=1000, hook_start=90, hook_end=0, coupler_start="short").cut_length(N5)
    with pytest.raises(GeometryError, match="extremo final termina en gancho"):
        HookedBar(a=1000, hook_start=0, hook_end=135, coupler_end="long", style="stirrup").cut_length(N3)
    with pytest.raises(GeometryError, match="extremo final termina en gancho"):
        CustomBar(legs=(400, 300), angles=(90,), hook_end=90, coupler_end="short").cut_length(N5)
    # la pata corta de una L es su gancho; el otro extremo sí admite coupler
    assert LBar(1200, 300, coupler_start="short").cut_length(N5).couplers == ("short",)
    assert "coupler_end" not in [p.name for p in LBar.parameters()]
    assert not any(p.name.startswith("coupler") for p in UBar.parameters())  # sus dos extremos son ganchos
    assert HookedBar(a=1000, hook_start=90, hook_end=0, coupler_end="long").cut_length(N5).couplers == ("long",)


def test_coupler_is_drawn_as_a_rectangle_at_the_end_with_its_type():
    sketch = StraightBar(3000, coupler_start="short", coupler_end="long").sketch(N5)
    assert [(b.kind, b.text) for b in sketch.boxes] == [("short", "RC"), ("long", "RL")]
    assert sketch.boxes[0].at == (0.0, 0.0) and sketch.boxes[0].heading_deg == 180.0  # sale por la izquierda
    assert sketch.boxes[1].at == (1.0, 0.0) and sketch.boxes[1].heading_deg == 0.0
    assert StraightBar(3000).sketch(N5).boxes == ()
    from exporters.shape_images import render_sketch

    def teal(image):  # píxeles del relleno del coupler
        data = image.tobytes()
        return sum(1 for r, g, b, a in zip(data[0::4], data[1::4], data[2::4], data[3::4])
                   if a > 200 and abs(r - 200) < 12 and abs(g - 228) < 12)

    with_coupler = render_sketch(StraightBar(3000, coupler_end="short").sketch(N5), 640, 240)
    long_coupler = render_sketch(StraightBar(3000, coupler_end="long").sketch(N5), 640, 240)
    assert teal(render_sketch(StraightBar(3000).sketch(N5), 640, 240)) == 0
    assert teal(with_coupler) > 200 and teal(long_coupler) > 1.4 * teal(with_coupler)  # rosca larga: rectángulo más largo


def test_coupler_counts_by_bar_and_type_in_the_project():
    element = StructuralElement("C-1", ElementType.COLUMN, "Nivel 1", 6)
    element.add_group("A", N5, StraightBar(3000, coupler_end="short"), 4)
    element.add_group("B", N5, StraightBar(3000, coupler_start="long", coupler_end="short"), 2)
    element.add_group("C", N6, ZBar(coupler_start="long"), 1)
    element.add_group("D", N5, StraightBar(2000), 5)  # sin couplers
    result = BBSCalculator().calculate([element])
    assert [(c.rebar.designation, c.kind, c.count) for c in result.coupler_counts()] == [
        ("#5", "short", 36), ("#5", "long", 12), ("#6", "long", 6)]
    assert BBSCalculator().calculate([]).coupler_counts() == ()


def test_couplers_are_listed_in_the_summaries_and_the_header_legend(tmp_path):
    from openpyxl import load_workbook

    from exporters.excel_exporter import ExcelExporter

    element = StructuralElement("C-1", ElementType.COLUMN, "Nivel 1", 2)
    element.add_group("A", N5, StraightBar(3000, coupler_end="short"), 4)
    element.add_group("B", N5, StraightBar(3000, coupler_start="long"), 1)
    result = BBSCalculator().calculate([element])
    wb = load_workbook(ExcelExporter().export(result, tmp_path / "c.xlsx"))
    summary = wb["Resumen"]
    texts = [str(summary.cell(r, 1).value) for r in range(1, summary.max_row + 1)]
    title = next(i for i, t in enumerate(texts, start=1) if "Couplers por barra y tipo" in t)
    assert summary.cell(title + 2, 3).value == "Rosca corta (RC)" and summary.cell(title + 2, 4).value == 8
    assert summary.cell(title + 3, 3).value == "Rosca larga (RL)" and summary.cell(title + 3, 4).value == 2
    assert summary.cell(title + 4, 4).value == 10  # total
    assert "RC = coupler de rosca corta, RL = coupler de rosca larga" in wb["Planilla BBS"]["B4"].value
    plain = StructuralElement("C-2", ElementType.COLUMN, "Nivel 1", 1)
    plain.add_group("A", N5, StraightBar(3000), 4)
    wb = load_workbook(ExcelExporter().export(BBSCalculator().calculate([plain]), tmp_path / "p.xlsx"))
    assert not any("Couplers" in str(c.value) for row in wb["Resumen"].iter_rows() for c in row)
    assert "RC =" not in wb["Planilla BBS"]["B4"].value


@pytest.mark.parametrize("shape", [
    Stirrup(220, 420, hook_start=90, hook_end=135), Stirrup(300, 400, hook_start=135, hook_end=90, ext_end=140),
    TriStirrup(300, 300, 300, hook_start=90), TrapStirrup(400, 250, 300, hook_end=90),
    UStirrup(300, 400, hook_start=90), RingTie(400, lap=0, hook_start=135, hook_end=90),
    RingTie(400, lap=250, hook_start=135, hook_end=90), Spiral(400, 75, 2000, hook_start=90, hook_end=135),
])
def test_mixed_hook_labels_also_fit_inside_the_picture(shape):
    from exporters.shape_images import BAR_COLOR, render_sketch

    image = render_sketch(shape.sketch(N3), 320, 120)
    width, height = image.size
    border = [image.getpixel((x, y)) for x in range(width) for y in (0, height - 1)]
    border += [image.getpixel((x, y)) for y in range(height) for x in (0, width - 1)]
    assert not [p for p in border if p[3] > 60 and sum((a - b) ** 2 for a, b in zip(p[:3], BAR_COLOR[:3])) > 25**2]
