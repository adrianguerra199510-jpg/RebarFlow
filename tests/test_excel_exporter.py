import pytest
from openpyxl import load_workbook

from core.element import ElementType, StructuralElement
from core.rebar import get_rebar
from core.shape_codes import LBar, Stirrup, StraightBar, UBar
from engine.bbs_calculator import BBSCalculator
from exporters.excel_exporter import ExcelExporter


@pytest.fixture
def result():
    n3, n5 = get_rebar("#3"), get_rebar("#5")
    beam = StructuralElement("V-101", ElementType.BEAM, "Nivel 1", 4)
    beam.add_group("A1", n5, StraightBar(5600), 3)
    beam.add_group("A2", n5, UBar.with_standard_hooks(n5, 5500), 2)
    beam.add_group("E1", n3, Stirrup(220, 420), 38)
    beam.add_group("E2", n3, Stirrup(220, 420, ext_start=50, ext_end=50), 2)  # gancho corto → aviso
    column = StructuralElement("C-1", ElementType.COLUMN, "Nivel 2", 6)
    column.add_group("C1", n5, LBar(1200, 300), 8)
    return BBSCalculator().calculate([beam, column])


def find_row(ws, text, column=1):
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, column).value == text)


def test_workbook_structure_and_totals(result, tmp_path):
    path = ExcelExporter(project_name="Proyecto X").export(result, tmp_path / "out" / "bbs.xlsx")
    assert path.exists()
    wb = load_workbook(path)
    assert wb.sheetnames == ["Planilla BBS", "Resumen", "Plan de corte"]

    ws = wb["Planilla BBS"]
    headers = [c.value for c in ws[7]]
    assert headers[:3] == ["Marca", "Barra", "Esquema"]
    # sin diámetro, forma, ganchos, observaciones ni columnas de cotas: las medidas van en el dibujo
    for removed in ("Ø (mm)", "Forma", "Ganchos (° · mm)", "Observaciones", "A", "B", "C", "D", "E"):
        assert removed not in headers
    assert headers == ["Marca", "Barra", "Esquema", "Long. corte (mm)", "Cant. por elem.", "Total barras",
                       "Peso unit. (kg)", "Long. total (m)", "Peso total (kg)"]
    weight_col = headers.index("Peso total (kg)") + 1

    total_row = find_row(ws, "TOTAL GENERAL")
    assert ws.cell(total_row, weight_col).value == pytest.approx(result.total_weight_kg)
    beam_row = find_row(ws, "Subtotal V-101")
    beam_weight = next(s for s in result.element_schedules if s.element.name == "V-101").total_weight_kg
    assert ws.cell(beam_row, weight_col).value == pytest.approx(beam_weight)
    assert ws.freeze_panes == "A8"


def test_line_values_and_warnings(result, tmp_path):
    ws = load_workbook(ExcelExporter().export(result, tmp_path / "bbs.xlsx"))["Planilla BBS"]
    headers = [c.value for c in ws[7]]
    col = {h: i + 1 for i, h in enumerate(headers)}
    stirrup = next(line for line in result.lines if line.mark == "E1")
    row = next(r for r in range(8, ws.max_row + 1) if ws.cell(r, 1).value == "E1")
    assert ws.cell(row, col["Long. corte (mm)"]).value == stirrup.cut_length_mm
    assert ws.cell(row, col["Total barras"]).value == 38 * 4

    # la marca con aviso va en rojo y el texto se lista en un bloque «Avisos» bajo la planilla
    short = next(r for r in range(8, ws.max_row + 1) if ws.cell(r, 1).value == "E2")
    assert ws.cell(short, 1).font.color.rgb.endswith("C00000") and ws.cell(row, 1).font.color is None
    assert ws.cell(find_row(ws, "Avisos"), 1).value == "Avisos"
    texts = [str(ws.cell(r, 1).value) for r in range(1, ws.max_row + 1) if "⚠" in str(ws.cell(r, 1).value)]
    assert len(texts) == 1 and "V-101 · marca E2" in texts[0] and "mínimo normativo" in texts[0]


def test_summary_sheet(result, tmp_path):
    ws = load_workbook(ExcelExporter().export(result, tmp_path / "bbs.xlsx"))["Resumen"]
    total_row = find_row(ws, "TOTAL")
    assert ws.cell(total_row, 7).value == pytest.approx(result.total_weight_kg)
    assert ws.cell(find_row(ws, "#3"), 7).value == pytest.approx(
        next(s.total_weight_kg for s in result.diameter_summaries if s.rebar.designation == "#3")
    )
    labels = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
    assert sum(1 for v in labels if v == "TOTAL") == 3  # tres tablas
    assert any("Sobrante de corte" in str(v) for v in labels)


def test_empty_result_exports(tmp_path):
    path = ExcelExporter().export(BBSCalculator().calculate([]), tmp_path / "empty.xlsx")
    assert load_workbook(path).sheetnames == ["Planilla BBS", "Resumen", "Plan de corte"]


def test_shape_images_are_embedded_one_per_line(result, tmp_path):
    ws = load_workbook(ExcelExporter().export(result, tmp_path / "bbs.xlsx"))["Planilla BBS"]
    assert len(ws._images) == len(result.lines)
    assert all(img.anchor._from.col == 2 for img in ws._images)  # columna "Esquema"


def test_images_can_be_disabled(result, tmp_path):
    path = ExcelExporter(include_shape_images=False).export(result, tmp_path / "plain.xlsx")
    ws = load_workbook(path)["Planilla BBS"]
    assert ws._images == []
    assert "Esquema" not in [c.value for c in ws[7]]
