import pytest
from openpyxl import load_workbook
from PIL import Image
from io import BytesIO

from core.element import ElementType, StructuralElement
from core.rebar import get_rebar
from core.shape_codes import Stirrup135, StraightBar, UBar
from core.shape_templates import Spiral
from engine.bbs_calculator import BBSCalculator
from engine.cutting_optimizer import CutPattern
from engine.cutting_report import plan_rows, plan_title, stock_comparison
from exporters.cutting_images import render_pattern_image, render_pattern_png
from exporters.excel_exporter import ExcelExporter
from exporters.pdf_exporter import PdfExporter
from sample_project import build_project

N3, N4, N5, N6 = (get_rebar(n) for n in ("#3", "#4", "#5", "#6"))


def sample_result(reusable=1000.0):
    project = build_project()
    return BBSCalculator(project.stock_length_mm, project.length_rounding_mm, reusable).calculate(project.elements)


def test_summary_carries_plan_marks_and_remnant_split():
    result = sample_result()
    by_bar = {s.rebar.designation: s for s in result.diameter_summaries}
    n5 = by_bar["#5"]
    assert n5.plan.stock_bars == n5.stock_bars == 28
    assert n5.waste_length_m == pytest.approx(n5.reusable_length_m + n5.scrap_length_m)
    assert n5.reusable_length_m == pytest.approx(27.52)  # retales de 1.60 m y 9.92 m
    assert n5.scrap_length_m == pytest.approx(0.8)
    assert set(n5.marks_for(2080.0)) == {"Z-1/Z1", "Z-1/Z2"}
    assert n5.marks_for(2080.4) == n5.marks_for(2080.0)  # tolerancia de 1 mm
    assert n5.total_weight_kg + n5.waste_weight_kg == pytest.approx(n5.purchase_weight_kg)
    assert result.total_scrap_weight_kg + result.total_reusable_weight_kg == pytest.approx(result.total_waste_weight_kg)


def test_reusable_threshold_zero_means_everything_is_scrap():
    result = sample_result(reusable=0.0)
    assert all(s.reusable_length_m == 0 for s in result.diameter_summaries)
    assert result.total_scrap_weight_kg == pytest.approx(result.total_waste_weight_kg)


def test_plan_beats_the_previous_first_fit_on_the_sample():
    """#5 necesitaba 30 barras con First Fit; el plan optimizado usa 28."""
    result = sample_result()
    assert {s.rebar.designation: s.stock_bars for s in result.diameter_summaries}["#5"] == 28
    assert result.total_waste_pct < 0.056  # antes 5.6 %


def test_spirals_come_from_coil_and_do_not_use_stock_bars():
    element = StructuralElement("C-9", ElementType.COLUMN, "Nivel 1", 2)
    element.add_group("E1", N4, Stirrup135(320, 320), 10)
    element.add_group("ES1", N4, Spiral(350, 75, 2000), 1)  # ~ 30 m: no cabe en una barra de 12 m
    result = BBSCalculator().calculate([element])
    (summary,) = result.diameter_summaries
    spiral_len = next(line for line in result.lines if line.mark == "ES1").total_length_m
    assert summary.coil_length_m == pytest.approx(spiral_len) and spiral_len > 12
    assert list(summary.pieces) == [next(line for line in result.lines if line.mark == "E1").cut_length_mm]
    assert summary.total_weight_kg == pytest.approx(sum(line.total_weight_kg for line in result.lines))
    assert summary.waste_length_m == pytest.approx(summary.plan.remnant_length_mm / 1000)  # el rollo no genera retales
    assert summary.purchased_length_m == pytest.approx(summary.plan.purchased_length_mm / 1000 + spiral_len)
    rows = plan_rows(summary)
    assert rows and "C-9/E1" in rows[0].marks_text
    assert "en rollo" in plan_title(summary)


def test_only_spirals_means_no_plan_rows():
    element = StructuralElement("P-1", ElementType.OTHER, "Nivel 1", 1)
    element.add_group("ES1", N4, Spiral(350, 75, 2000), 1)
    (summary,) = BBSCalculator().calculate([element]).diameter_summaries
    assert plan_rows(summary) == [] and summary.stock_bars == 0
    assert summary.waste_pct == 0 and summary.purchased_length_m == pytest.approx(summary.coil_length_m)


def test_plan_rows_describe_the_sequence():
    summary = next(s for s in sample_result().diameter_summaries if s.rebar.designation == "#5")
    rows = plan_rows(summary)
    assert [r.step for r in rows] == [1, 2, 3]
    assert sum(r.count for r in rows) == summary.stock_bars
    assert rows[0].bars_text == "1–16" and rows[-1].bars_text == "28"
    assert rows[0].cuts_text == "4 × 2,080 + 2 × 1,815"
    assert [r.remnant_use for r in rows] == ["Desperdicio", "Reutilizable", "Reutilizable"]
    title = plan_title(summary)
    assert "#5" in title and "28 barras de 12.00 m" in title and "plan lineal" in title


def test_stock_comparison_recommends_the_length_with_least_purchase():
    result = sample_result()
    comparisons = stock_comparison(result)
    assert [c.summary.rebar.designation for c in comparisons] == ["#3", "#5", "#6", "#7"]
    for c in comparisons:
        assert c.best_mm == min(c.options, key=lambda o: (round(o.purchased_m, 6), -o.stock_length_mm)).stock_length_mm
    n6 = next(c for c in comparisons if c.summary.rebar.designation == "#6")  # piezas de 6 m: no caben en 6 m
    assert {o.stock_length_mm for o in n6.options} == {9000.0, 12000.0}


def test_pattern_images_have_the_requested_size_and_colour_remnants():
    pattern = CutPattern((3900.0, 3900.0, 3900.0), 1, 12000.0)
    png = render_pattern_png(pattern, 840, 60, 1000.0)
    assert Image.open(BytesIO(png)).size == (840, 60)
    reusable = CutPattern((7000.0, 4000.0), 1, 12000.0)  # retal de 1 m
    green = render_pattern_image(reusable, 600, 40, 1000.0).convert("RGB").getpixel((592, 20))
    gray = render_pattern_image(reusable, 600, 40, 5000.0).convert("RGB").getpixel((592, 20))
    assert green != gray and green[1] > green[0] and abs(gray[0] - gray[2]) < 12


def test_excel_plan_sheet(tmp_path):
    result = sample_result()
    wb = load_workbook(ExcelExporter(project_name="X").export(result, tmp_path / "p.xlsx"))
    ws = wb["Plan de corte"]
    total_steps = sum(len(plan_rows(s)) for s in result.diameter_summaries)
    assert len(ws._images) == total_steps
    texts = [str(ws.cell(r, 1).value) for r in range(1, ws.max_row + 1)]
    assert any(t.startswith("#5 · Grado 60") for t in texts)
    assert any("4 × 2,080 + 2 × 1,815" == str(ws.cell(r, 5).value) for r in range(1, ws.max_row + 1))
    summary = wb["Resumen"]
    labels = [str(summary.cell(r, 1).value) for r in range(1, summary.max_row + 1)]
    assert any("Longitud comercial recomendada" in t for t in labels)
    assert any(summary.cell(r, 10).value is not None and summary.cell(r, 1).value == "TOTAL" for r in range(1, summary.max_row + 1))


def test_excel_hook_column_matches_the_schedule(tmp_path):
    result = sample_result()
    ws = load_workbook(ExcelExporter().export(result, tmp_path / "h.xlsx"))["Planilla BBS"]
    headers = [c.value for c in ws[7]]
    column = headers.index("Ganchos (° · mm)") + 1
    values = {ws.cell(r, 1).value: ws.cell(r, column).value for r in range(8, ws.max_row + 1) if ws.cell(r, 1).value}
    assert values["E1"] == "2 × 135° · 76" or values["E1"] == "2 × 135° · 75"  # estribo #3 (75) o #4 (76)
    assert values["A2"].startswith("2 × 90°")  # U con ganchos de 90°
    assert values["A1"] in (None, "")  # barra recta: sin ganchos


def test_pdf_has_the_cutting_plan_pages(tmp_path):
    path = PdfExporter().export(sample_result(), tmp_path / "p.pdf")
    data = path.read_bytes()
    assert data.startswith(b"%PDF") and path.stat().st_size > 20_000
    import re
    assert len(re.findall(rb"/Type /Page\b(?!s)", data)) >= 4  # planilla + resumen + plan


def test_empty_result_exports_with_the_plan_sheet(tmp_path):
    empty = BBSCalculator().calculate([])
    assert load_workbook(ExcelExporter().export(empty, tmp_path / "e.xlsx")).sheetnames[-1] == "Plan de corte"
    assert PdfExporter().export(empty, tmp_path / "e.pdf").read_bytes().startswith(b"%PDF")


def test_catalog_project_uses_every_builtin_template_and_exports(tmp_path):
    from sample_project import build_catalog_project

    project = build_catalog_project()
    element = project.elements[0]
    assert element.element_type is ElementType.OTHER
    assert {g.shape.code for g in element.groups} == {"R", "L", "U", "E", "E90", "G", "EU", "Z", "Q", "O", "ET", "ETR", "AR", "ES", "X"}
    result = BBSCalculator(project.stock_length_mm, project.length_rounding_mm, project.reusable_remnant_mm).calculate(project.elements)
    assert len(result.lines) == 16 and result.total_weight_kg > 0
    assert not [w for line in result.lines for w in line.warnings]  # el ejemplo cumple los mínimos de la norma
    assert next(line for line in result.lines if line.mark == "ES1").shape.from_coil
    assert ExcelExporter(project.name).export(result, tmp_path / "c.xlsx").exists()
    assert PdfExporter(project.name).export(result, tmp_path / "c.pdf").read_bytes().startswith(b"%PDF")


def test_images_label_hook_lengths_only_when_the_bar_is_known():
    from core.shape_codes import LBar
    from exporters.shape_images import render_shape_png

    def amber_pixels(data):
        image = Image.open(BytesIO(data)).convert("RGBA")
        pixels = image.load()
        count = 0
        for x in range(image.width):
            for y in range(image.height):
                r, g, b, a = pixels[x, y]
                count += a > 120 and r > 150 and 60 < g < 120 and b < 40  # ámbar de los rótulos de gancho
        return count

    shape = LBar.with_standard_hook(N4, a=1000)
    assert amber_pixels(render_shape_png(shape, 640, 240, None)) == 0
    assert amber_pixels(render_shape_png(shape, 640, 240, N4)) > 40  # aparece «G 152» en ámbar
