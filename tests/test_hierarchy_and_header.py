"""Obras sin niveles (un puente): jerarquía opcional, datos de encabezado de la planilla y logo."""

import base64
import shutil
import subprocess
from io import BytesIO

import pytest
from openpyxl import load_workbook
from PIL import Image

from core.element import ElementType, StructuralElement
from core.project import Project
from core.rebar import get_rebar
from core.serialization import load_project, project_from_dict, project_to_dict, save_project
from core.shape_codes import StraightBar
from engine.bbs_calculator import BBSCalculator
from exporters.excel_exporter import ExcelExporter
from exporters.logo import LogoError, fitted_size, normalize_logo
from exporters.pdf_exporter import PdfExporter
from sample_project import build_bridge_project, build_project

N4 = get_rebar("#4")


def image_bytes(size=(400, 120), color=(20, 100, 110), fmt="JPEG") -> bytes:
    buffer = BytesIO()
    Image.new("RGB", size, color).save(buffer, format=fmt)
    return buffer.getvalue()


# ------------------------------------------------------------------ modelo
def test_element_levels_are_optional_and_form_a_location():
    bare = StructuralElement("Barrera", ElementType.WALL)
    assert bare.level == bare.substructure == "" and bare.path == () and bare.location == ""
    assert bare.heading == "Barrera  ·  Muro  ·  1 und"
    full = StructuralElement("Losa 1", ElementType.SLAB, " C9.6-1 ", 2, substructure="GENERICA ")
    assert full.path == ("C9.6-1", "GENERICA") and full.location == "C9.6-1 / GENERICA"
    assert full.heading == "Losa 1  ·  Losa  ·  C9.6-1 / GENERICA  ·  2 und"
    only_sub = StructuralElement("Viga", ElementType.BEAM, substructure="Pila")
    assert only_sub.path == ("Pila",)
    with pytest.raises(ValueError, match="nombre"):
        StructuralElement("  ", ElementType.BEAM)


def test_project_groups_header_fields_and_label():
    project = build_bridge_project()
    assert project.group_label == "Estructura" and project.levels == ["E-1", "P-1"]
    assert project.substructures == ["GENERICA", "Cimentación", "Pila"]  # el elemento sin grupos no aporta vacíos
    assert [label for label, _ in project.header_fields()] == [
        "Cliente", "Cód. obra", "N° planilla", "Plano", "P. global", "F. entrega"]
    assert Project().header_fields() == () and Project(group_label="  ").group_label == "Nivel"
    assert dict(Project(client=" ACME ", drawing="P-1").header_fields()) == {"Cliente": "ACME", "Plano": "P-1"}


def test_result_groups_and_titles():
    bridge = build_bridge_project()
    result = BBSCalculator(group_label=bridge.group_label).calculate(bridge.elements)
    assert result.levels == ("E-1", "P-1", "") and result.has_level_groups
    assert result.level_title("P-1") == "P-1" and result.level_title("") == "Sin estructura"
    assert sum(sum(row.values()) for row in result.weight_by_level_and_diameter().values()) == pytest.approx(result.total_weight_kg)
    flat = StructuralElement("Muro", ElementType.WALL)
    flat.add_group("A1", N4, StraightBar(3000), 4)
    single = BBSCalculator().calculate([flat])
    assert single.levels == ("",) and not single.has_level_groups and single.level_title("") == "Sin nivel"


def test_same_name_is_allowed_in_different_substructures_but_not_twice_in_the_same():
    a = StructuralElement("Losa", ElementType.SLAB, "E-1", substructure="Norte")
    b = StructuralElement("Losa", ElementType.SLAB, "E-1", substructure="Sur")
    BBSCalculator().calculate([a, b])
    with pytest.raises(ValueError, match="repetido en 'E-1 / Norte'"):
        BBSCalculator().calculate([a, StructuralElement("Losa", ElementType.SLAB, "E-1", substructure="Norte")])
    with pytest.raises(ValueError, match="repetido"):
        BBSCalculator().calculate([StructuralElement("Losa", ElementType.SLAB), StructuralElement("Losa", ElementType.SLAB)])


# ------------------------------------------------------------------ archivo
def test_bridge_project_round_trips_with_header_and_logo(tmp_path):
    project = build_bridge_project()
    project.logo_png = normalize_logo(image_bytes())
    loaded = load_project(save_project(project, tmp_path / "puente.rebarflow"))
    assert loaded.group_label == "Estructura" and loaded.header_fields() == project.header_fields()
    assert loaded.logo_png == project.logo_png and loaded.author == project.author
    assert [(e.name, e.level, e.substructure) for e in loaded.elements] == [
        (e.name, e.level, e.substructure) for e in project.elements]
    data = project_to_dict(project)
    assert base64.b64decode(data["logo"]) == project.logo_png and data["header"]["work_code"] == "00001"


def test_old_project_files_without_the_new_fields_still_open():
    data = project_to_dict(build_project())
    for key in ("group_label", "header", "logo"):
        data.pop(key)
    for element in data["elements"]:
        element.pop("substructure")
    loaded = project_from_dict(data)
    assert loaded.group_label == "Nivel" and loaded.header_fields() == () and loaded.logo_png == b""
    assert loaded.levels == ["Cimentación", "Nivel 1", "Nivel 2"]


# ------------------------------------------------------------------ logo
def test_logo_is_normalized_to_a_small_png():
    big = image_bytes((3000, 1000))
    png = normalize_logo(big)
    image = Image.open(BytesIO(png))
    assert image.format == "PNG" and image.mode == "RGBA" and image.size[0] <= 640 and image.size[1] <= 240
    assert image.size[0] / image.size[1] == pytest.approx(3.0, rel=0.02)  # no se deforma
    assert Image.open(BytesIO(normalize_logo(image_bytes((50, 20), fmt="BMP")))).size == (50, 20)  # no se amplía
    with pytest.raises(LogoError, match="No se pudo leer"):
        normalize_logo(b"esto no es una imagen")
    width, height = fitted_size(png, 100, 100)
    assert width == pytest.approx(100) and height == pytest.approx(100 / 3, rel=0.02)


# ------------------------------------------------------------------ exportación
def bridge_result():
    project = build_bridge_project()
    return project, BBSCalculator(group_label=project.group_label).calculate(project.elements)


def test_excel_header_shows_the_work_data_and_the_logo(tmp_path):
    project, result = bridge_result()
    logo = normalize_logo(image_bytes())
    plain = load_workbook(ExcelExporter(project.name, project.author).export(result, tmp_path / "a.xlsx"))["Planilla BBS"]
    ws = load_workbook(ExcelExporter(project.name, project.author, info=project.header_fields(), logo_png=logo)
                       .export(result, tmp_path / "b.xlsx"))["Planilla BBS"]
    assert [c.value for c in plain[7]][0] == "Marca"  # sin datos ni logo: el mismo diseño de siempre
    texts = [str(ws.cell(r, 2).value) for r in range(1, 12)]
    data_row = next(t for t in texts if "Cliente: Consorcio de ejemplo" in t)
    assert "Cód. obra: 00001" in data_row and "N° planilla: 1.001" in data_row
    assert any("Plano: S-101" in t and "F. entrega: 15/10/2026" in t for t in texts)
    header_row = next(r for r in range(1, 20) if ws.cell(r, 1).value == "Marca")
    assert header_row > 7 and ws.freeze_panes == f"A{header_row + 1}" and ws.print_title_rows == f"${header_row}:${header_row}"
    logos = [i for i in ws._images if i.anchor._from.row == 0]
    assert len(logos) == 1 and logos[0].anchor._from.col == 0  # arriba a la izquierda
    band = next(str(ws.cell(r, 1).value) for r in range(header_row, ws.max_row) if "Losa de transición 1" in str(ws.cell(r, 1).value))
    assert band == "Losa de transición 1  ·  Losa  ·  E-1 / GENERICA  ·  2 und"
    assert any(str(ws.cell(r, 1).value) == "Barrera de concreto  ·  Muro  ·  4 und" for r in range(header_row, ws.max_row))
    sheets = load_workbook(tmp_path / "b.xlsx")  # el resumen y el plan de corte se desplazan con el encabezado
    assert any(str(c.value) == "Barra" for row in sheets["Resumen"].iter_rows() for c in row)


def test_summary_has_the_consolidated_table_only_with_two_or_more_groups(tmp_path):
    project, result = bridge_result()
    titles = lambda ws: [str(ws.cell(r, 1).value) for r in range(1, ws.max_row + 1)]  # noqa: E731
    bridge_sheet = load_workbook(ExcelExporter().export(result, tmp_path / "p.xlsx"))["Resumen"]
    assert any("Consolidado por estructura" in t for t in titles(bridge_sheet))
    header = next(r for r in range(1, bridge_sheet.max_row + 1) if "Consolidado por estructura" in str(bridge_sheet.cell(r, 1).value)) + 1
    assert [bridge_sheet.cell(header, c).value for c in range(1, 6)] == ["Barra", "E-1", "P-1", "Sin estructura", "Total"]
    flat = StructuralElement("Muro", ElementType.WALL)
    flat.add_group("A1", N4, StraightBar(3000), 4)
    flat_sheet = load_workbook(ExcelExporter().export(BBSCalculator().calculate([flat]), tmp_path / "f.xlsx"))["Resumen"]
    assert not any("Consolidado" in t for t in titles(flat_sheet))
    assert [t[:2] for t in titles(flat_sheet) if t[:2] in ("1.", "2.", "3.", "4.")] == ["1.", "2.", "3."]  # numeración sin huecos


@pytest.mark.skipif(shutil.which("pdftotext") is None, reason="requiere pdftotext (poppler)")
def test_pdf_header_logo_location_and_summary(tmp_path):
    project, result = bridge_result()
    logo = normalize_logo(image_bytes())
    path = PdfExporter(project.name, project.author, info=project.header_fields(), logo_png=logo).export(result, tmp_path / "p.pdf")
    text = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True, check=True).stdout
    assert "Cliente: Consorcio de ejemplo" in text and "Cód. obra: 00001" in text and "Plano: S-101" in text
    assert "E-1 / GENERICA" in text and "Barrera de concreto" in text
    assert "Consolidado por estructura" in text and "Sin estructura" in text
    assert b"/Subtype /Image" in PdfExporter(logo_png=logo).export(result, tmp_path / "l.pdf").read_bytes()
    flat = StructuralElement("Muro", ElementType.WALL)
    flat.add_group("A1", N4, StraightBar(3000), 4)
    flat_pdf = PdfExporter().export(BBSCalculator().calculate([flat]), tmp_path / "f.pdf")
    assert "Consolidado" not in subprocess.run(["pdftotext", "-layout", str(flat_pdf), "-"], capture_output=True, text=True).stdout
